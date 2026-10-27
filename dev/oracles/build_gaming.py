"""Build decks a human would call obviously wrong for the keybank reference.

G1 scramble : keybank.pptx, every slide-level shape moved to a random position (sizes kept)
G2 stack    : keybank.pptx, every slide-level shape moved to (0,0): everything overlaps
G3 overflow : the honest generated section with every run's text repeated 12x (every box overflows)
G4 giant    : keybank template, 12 slides, one page-sized text shape of Arial 12 black per slide
G5 random   : keybank template, 12 slides, 9 random boxes per slide at random positions/sizes
              carrying keybank's values (Arial 12 black, 113D63 / CCD1D7 fills, 24pt title)
"""
import copy
import random
import shutil
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": A, "p": P}
W, H = 9144000, 6858000

out_dir = Path(sys.argv[1])
keybank = sys.argv[2]
honest = sys.argv[3]
template = sys.argv[4]
out_dir.mkdir(parents=True, exist_ok=True)
rng = random.Random(7)


def rewrite_slides(src: str, dst: str, edit) -> dict:
    stats = {}
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("ppt/slides/slide") and item.filename.endswith(".xml"):
                root = etree.fromstring(data)
                stats[item.filename] = edit(root)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    return stats


def _xfrms(root):
    tree = root.find(f"{{{P}}}cSld/{{{P}}}spTree")
    for child in tree:
        tag = etree.QName(child).localname
        if tag == "graphicFrame":
            x = child.find(f"{{{P}}}xfrm")
        elif tag == "grpSp":
            x = child.find(f"{{{P}}}grpSpPr/{{{A}}}xfrm")
        elif tag in ("sp", "pic", "cxnSp"):
            x = child.find(f"{{{P}}}spPr/{{{A}}}xfrm")
        else:
            continue
        if x is None:
            continue
        off = x.find(f"{{{A}}}off")
        ext = x.find(f"{{{A}}}ext")
        if off is None or ext is None:
            continue
        yield off, ext


def scramble(root):
    n = 0
    for off, ext in _xfrms(root):
        cx, cy = int(ext.get("cx")), int(ext.get("cy"))
        off.set("x", str(rng.randint(0, max(0, W - cx))))
        off.set("y", str(rng.randint(0, max(0, H - cy))))
        n += 1
    return n


def stack(root):
    n = 0
    for off, ext in _xfrms(root):
        off.set("x", "0")
        off.set("y", "0")
        n += 1
    return n


def overflow(root):
    n = 0
    for t in root.iter(f"{{{A}}}t"):
        if t.text and t.text.strip():
            t.text = (t.text + " ") * 12
            n += 1
    return n


rewrite_slides(keybank, str(out_dir / "g1_scramble.pptx"), scramble)
rewrite_slides(keybank, str(out_dir / "g2_stack.pptx"), stack)
rewrite_slides(honest, str(out_dir / "g3_overflow.pptx"), overflow)

LOREM = ("Pro forma tangible book value per share accretion dilution at announcement and at close "
         "assuming 35 percent cost synergies phased in over two years and a 1.0x exchange ratio. ")


def style_run(run, size=12, color="000000", bold=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text_shape(slide, x, y, w, h, text, size=12, color="000000", fill=None, align=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shp.fill.background()
    shp.line.fill.background()
    tf = shp.text_frame
    tf.word_wrap = True
    paras = text if isinstance(text, list) else [text]
    for i, ptxt in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = ptxt
        style_run(r, size, color)
    return shp


def giant():
    prs = Presentation(template)
    layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
    for i in range(12):
        s = prs.slides.add_slide(layout)
        for ph in list(s.placeholders):
            if ph.placeholder_format.type is not None and "TITLE" in str(ph.placeholder_format.type):
                ph.text_frame.text = f"Giant box page {i + 1}"
                for p in ph.text_frame.paragraphs:
                    for r in p.runs:
                        style_run(r, 24, "000000")
            else:
                ph._element.getparent().remove(ph._element)
        add_text_shape(s, 0, 0, W, H, [LOREM * 3] * 14)
    prs.save(str(out_dir / "g4_giant.pptx"))


def random_boxes():
    prs = Presentation(template)
    layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
    for i in range(12):
        s = prs.slides.add_slide(layout)
        for ph in list(s.placeholders):
            if ph.placeholder_format.type is not None and "TITLE" in str(ph.placeholder_format.type):
                ph.text_frame.text = f"Random boxes page {i + 1}"
                for p in ph.text_frame.paragraphs:
                    for r in p.runs:
                        style_run(r, 24, "000000")
            else:
                ph._element.getparent().remove(ph._element)
        for j in range(9):
            w = rng.randint(W // 12, W // 2)
            h = rng.randint(H // 14, H // 3)
            x = rng.randint(0, W - w)
            y = rng.randint(0, H - h)
            roll = rng.random()
            if roll < 0.15:
                add_text_shape(s, x, y, w, h, LOREM[: rng.randint(20, 90)], 12, "FFFFFF", fill="113D63")
            elif roll < 0.3:
                add_text_shape(s, x, y, w, h, LOREM[: rng.randint(20, 90)], 12, "000000", fill="CCD1D7")
            else:
                add_text_shape(s, x, y, w, h, LOREM[: rng.randint(40, 160)], 12, "000000")
    prs.save(str(out_dir / "g5_random.pptx"))


giant()
random_boxes()
print("built", sorted(p.name for p in out_dir.glob("g*.pptx")))
