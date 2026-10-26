#!/usr/bin/env python3
"""Write slides onto the reference's own package from a resolved slide spec and a style guide.

    emit_pptx.py resolved.json --template tpl.pptx --guide guide.json -o out.pptx [--chrome auto|none]

`tpl.pptx` is the reference with its slides removed (pptx_tools.py template): masters,
layouts, theme and chrome stay, so every inherited value is the reference's (master
reuse: PowerPoint opens and re-saves such a package without repair). The spec carries TOKEN ROLES and content only. A colour is a role name in
the guide (`color.ink`, `color.accent2`, `color.extra_3`), a text style is a type role
(`type.body`); a literal hex or point size in the spec is refused unless the element is
marked `"foreign": true`. Boxes are integer EMU from the layout solver.

Spec shape (one slide shown):
{
 "slides": [
  {"layout": "Title and Content",
   "title": {"text": "Action title", "role": "type.title"},
   "elements": [
    {"id": "body", "kind": "text", "box": [x, y, w, h],
     "paragraphs": [{"role": "type.body", "bullet_level": 0,
                     "runs": [{"text": "Plain "}, {"text": "bold", "emphasis": true}],
                     "lines": ["Plain bold"]}],
     "fill_role": null, "line_role": null, "anchor": "t"},
    {"id": "comps", "kind": "table", "box": [...], "header_rows": 1,
     "columns": [{"key": "name", "label": "Company", "value_kind": "label"}, {"key": "ev", "label": "EV ($m)", "value_kind": "number"}],
     "rows": [["Acme", "1,234"]], "column_widths_emu": [...], "row_heights_emu": [...], "totals_last_row": false,
     "header_role": "type.table_header", "cell_role": "type.table_cell"},
    {"id": "chart", "kind": "chart", "box": [...], "type": "COLUMN_STACKED", "categories": ["2023", "2024"],
     "series": [{"name": "Revenue", "values": [1.0, 2.0], "fill_role": "color.accent2", "line_role": "none",
                 "labels": {"show": true, "position": "center", "number_format": "0.0", "role": "type.data_label", "ink_role": "color.ink_inverse"},
                 "points": [{"index": 1, "fill_role": "color.accent1", "label": {"ink_role": "color.ink"}}]}],
     "axes": {"category": {"visible": true, "line_role": "color.rule", "line_width_pt": 0.75, "tick_marks": "none", "reverse": false,
                           "gridlines": "none", "labels": {"show": true, "role": "type.axis_label", "ink_role": "color.ink_muted"}},
              "value": {"visible": false}},
     "gap_width": 50, "overlap": 100, "legend": "none", "text": {...}, "plot_fill_role": "none", "title": "none"},
    {"id": "rule", "kind": "rule", "box": [...], "line_role": "color.rule"},
    {"id": "band", "kind": "rect", "box": [...], "fill_role": "color.accent1", "prst": "rect"},
    {"id": "logo", "kind": "image", "box": [...], "path": "logo.png"}
   ],
   "notes": "speaker notes"}
 ]
}
Text bodies are written with zero insets, square wrap and explicit noAutofit, paragraph
spacing and line pitch as absolute points, and, when the solver measured the wrap, the
measured lines as literal breaks: the file carries the wrap the fit decision was made on.
Sizes come from the same place: a paragraph's `fit` (the solver's resolved entry, sz,
lnSpc_spcPts, spcBef, spcAft, spc in hundredths) and a table's `fit_header`, `fit_cell`
and `fit_padding_emu` are what the file is written at, so a page fitted at scale 0.89
is written at 0.89. Measured lines without `fit`, or fitted row heights without
`fit_cell`, are refused: lines measured at one size and written at another overflow.

A chart is the spec the model wrote: `type` is any XL_CHART_TYPE member python-pptx
writes, every colour a role or "none", every label, axis, legend and spacing choice
explicit. A choice the spec omits is written as the neutral value in CHART_DEFAULTS and
named in the slide note's `defaults_taken`. Waterfalls, football fields and mekkos are
recipes over this spec (references/chart-recipes.md), not code paths. No ink is
chosen here for any fill: tables and charts write the ink the spec or guide names, and
the build's contrast gate refuses a pair below 3:1.

What the file carries so that PowerPoint and LibreOffice draw the same page (measured on
both): both cells of every shared table edge declare the edge identically, every chart
plot group writes c:varyColors val=0 (the one chart value written unasked), no rule
is written below 0.5pt unless the guide's own rule width says so. Every generated slide carries a tags part (p:custDataLst/p:tags,
one p:tag name="pptx-skill" val="generated"), which survives a PowerPoint save
(measured): `pptx_tools.py text` reads it back as the slide's origin.
"""

from __future__ import annotations

import argparse
import copy
import json
import math
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.chart.datalabel import DataLabels
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LABEL_POSITION,
    XL_LEGEND_POSITION,
    XL_MARKER_STYLE,
    XL_TICK_LABEL_POSITION,
    XL_TICK_MARK,
)
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.dont_write_bytecode = True  # the skill directory is never written, not even a bytecode cache
sys.path.insert(0, str(Path(__file__).resolve().parent))
from chart_geometry import chart_family  # noqa: E402

# the solver's own rules for case and header cells: it measured this text
from layout import (  # noqa: E402
    DEFAULT_CELL_INSET_EMU,
    apply_case,
    apply_case_runs,
    header_labels,
    header_spans,
)

EMU_PER_PT = 12700
HEX_RE = re.compile(r"^[0-9A-Fa-f]{6}$")
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
# The origin tag every generated slide carries (one p:tag in its own tags part). deckgen
# 2E-1 measured the part coming back out of a PowerPoint save with the p:tag intact.
ORIGIN_TAG_NAME = "pptx-skill"
ORIGIN_TAG_VALUE = "generated"
# The thinnest stroke PowerPoint still inks: at 0.25pt it lays 0.48px where LibreOffice
# lays 1.27px (deckgen E38 class 5); at 0.5pt both are above one pixel.
RULE_WEIGHT_FLOOR_EMU = 6350
# Chart plot groups whose schema carries c:varyColors (ECMA-376 21.2.2). The element
# defaults to TRUE and the two renderers read the default oppositely (E38 class 10).
VARY_COLORS_GROUPS = {"barChart", "bar3DChart", "lineChart", "line3DChart", "pieChart", "pie3DChart", "ofPieChart", "doughnutChart",
                      "areaChart", "area3DChart", "scatterChart", "radarChart", "bubbleChart"}
# Children that precede c:varyColors in those groups' sequences.
_VARY_COLORS_PREDECESSORS = {"barDir", "grouping", "scatterStyle", "radarStyle", "ofPieType"}


class SpecError(SystemExit):
    pass


# --- guide access -------------------------------------------------------------------------------


class Guide:
    def __init__(self, data: dict) -> None:
        self.g = data
        self.roles = data.get("typography", {}).get("roles", {})
        self.families = data.get("typography", {}).get("families", {})
        colors = dict(data.get("color", {}).get("roles", {}))
        colors.update(data.get("extra_roles", {}))
        self.colors = colors
        self.spacing = data.get("spacing", {}).get("steps", {})
        self.tables = data.get("tables") or {}
        self.charts = data.get("charts") or {}
        self.text_container = (data.get("register_stats") or {}).get("text_container", "text_shape")
        #: (element id, what) for every value this emit wrote that neither the spec nor the
        #: guide states; emit() resets it per slide and reports it as `defaults_taken`
        self.taken: list[tuple[str, str]] = []

    def take(self, where: str, what: str, *, look: bool = False) -> None:
        """List a value written that neither the spec nor the guide states; a `look` is
        refused first where nobody chose it (unchosen)."""
        if look:
            self.unchosen(where, what)
        if (where, what) not in self.taken:
            self.taken.append((where, what))

    def unchosen(self, where: str, what: str) -> None:
        """A look (a face, size, weight, alignment, colour, fill, stroke or bar geometry) that
        neither the spec nor the guide states falls to the template's theme, a renderer's own
        style or a constant here. A guide made with no reference has no theme of anyone's to
        fall to (the blank template's is python-pptx's stock Office theme), so it is refused
        by name there; with a reference it is listed like any other default."""
        import plan
        if not plan.has_reference(self.g):
            raise SpecError(f"{where}: {what}: with no reference nothing is drawn in a look nobody chose (the template's theme, "
                            "a renderer's own style or a constant in the emitter); write this choice")

    def hex(self, role: str | None, *, where: str) -> str | None:
        if role is None:
            return None
        if HEX_RE.match(role):
            raise SpecError(f"{where}: literal colour {role!r}; use a colour role from the guide")
        entry = self.colors.get(role)
        if entry is None and role in (self.g.get("absent_roles") or {}):
            raise SpecError(f"{where}: colour role {role!r}: the completed guide records that the reference has none "
                            f"({self.g['absent_roles'][role].get('why')}); use a role the reference shows")
        if entry is None:
            raise SpecError(f"{where}: unknown colour role {role!r}; guide roles: {sorted(self.colors)}")
        return entry["hex"]

    def type_spec(self, role: str, *, where: str) -> dict:
        spec = self.roles.get(role)
        if spec is None and role in (self.g.get("absent_roles") or {}):
            raise SpecError(f"{where}: type role {role!r}: the completed guide records that the reference has none "
                            f"({self.g['absent_roles'][role].get('why')}); use a role the reference sets")
        if spec is None:
            raise SpecError(f"{where}: unknown type role {role!r}; guide roles: {sorted(self.roles)}")
        if spec.get("partial"):
            # a ladder-derived role has a size only: the other members ride on the body role
            base = dict(self.roles.get("type.body", {}))
            base.update({k: v for k, v in spec.items() if v is not None})
            spec = base
        return spec

    def face(self, spec: dict) -> str:
        """The face a type role is written in: its family, or the guide's face for its
        major/minor token. A role naming no face the guide can resolve is refused: a stand-in
        written here would show one face while the solver measured another."""
        fam = spec.get("family")
        face = (self.families.get(fam) or (self.families.get("minor") if fam == "major" else None)) if fam in ("major", "minor") else fam or self.families.get("minor")
        if not face:
            raise SpecError(f"a type role's family {fam!r} names no face and the guide's typography.families has no "
                            f"{'major or minor' if fam in ('major', 'minor') else 'minor'} face; complete /typography/families")
        return face

    def space(self, role: str | None, default: int = 0) -> int:
        if role is None:
            return default
        entry = self.spacing.get(role)
        if entry is None:
            raise SpecError(f"unknown spacing role {role!r}; guide steps: {sorted(self.spacing)}")
        return int(entry["emu"])


# --- xml helpers ----------------------------------------------------------------------------------


def _solid_fill(parent: etree._Element, hex_color: str) -> None:
    for child in list(parent):
        if child.tag in (qn("a:noFill"), qn("a:solidFill"), qn("a:gradFill"), qn("a:blipFill"), qn("a:pattFill")):
            parent.remove(child)
    sf = etree.SubElement(parent, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", hex_color.upper())


def _without_theme_style(shape):
    """python-pptx gives a new autoshape or connector a p:style whose lnRef and fillRef
    paint accent1 and whose effectRef paints a theme effect style (the blank template's
    is a drop shadow, which LibreOffice renders under every rule). The emitter writes
    each fill and line itself, so no theme reference stays on the shape."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    return shape


def _insert_fill_first(parent: etree._Element, fill: etree._Element) -> None:
    """DrawingML orders fill before ln/effects: put a fill right after xfrm/prstGeom."""
    anchor = 0
    for i, child in enumerate(parent):
        if child.tag in (qn("a:xfrm"), qn("a:prstGeom"), qn("a:custGeom")):
            anchor = i + 1
    parent.insert(anchor, fill)


def _set_line(sp_pr: etree._Element, hex_color: str | None, width_emu: int | None, dash: str | None = None) -> None:
    ln = sp_pr.find(qn("a:ln"))
    if ln is None:
        ln = etree.SubElement(sp_pr, qn("a:ln"))
    for child in list(ln):
        ln.remove(child)
    if hex_color is None:
        etree.SubElement(ln, qn("a:noFill"))
        return
    if width_emu:
        ln.set("w", str(int(width_emu)))
    sf = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", hex_color.upper())
    if dash and dash != "solid":
        etree.SubElement(ln, qn("a:prstDash")).set("val", dash)


ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}


def paragraph_align(para: dict, spec: dict) -> str:
    """The alignment a paragraph is written with: its own, else its role's, else left."""
    return para.get("align") or spec.get("alignment") or "left"


def _end_paragraph_like(p, run) -> None:
    """The paragraph's end mark set like its run (a:endParaRPr, a copy of the run's rPr).
    A paragraph whose run is empty (an empty table cell) is drawn at the end mark's size;
    without one LibreOffice draws it at its default 18pt, so the row grows past the
    solver's fitted height and every row below moves (measured: a 12pt keybank row drawn
    about 13 px taller at 1600 px, and the boxes below it landed on the next row). Measured
    on LibreOffice 24.2: the end mark beside an empty run, or the run removed with no end
    mark, still drew the tall row; an empty cell needs both, so write_table drops its run."""
    rpr = run._r.find(qn("a:rPr"))
    for old in p._p.findall(qn("a:endParaRPr")):
        p._p.remove(old)
    if rpr is None:
        return
    end = copy.deepcopy(rpr)
    end.tag = qn("a:endParaRPr")
    p._p.append(end)


def run_style(run: dict, where: str) -> tuple[bool | None, bool | None]:
    """(bold, italic) a run of a paragraph is set in: its own `bold` and `italic` (true or
    false), `emphasis: true` meaning bold; None where it says nothing (the role's own). The
    one reading the solver measures (build_deck.mark_measured_weights), this file writes and
    the contrast gate scores."""
    for key in ("emphasis", "bold", "italic"):
        if key in run and not isinstance(run[key], bool):
            raise SpecError(f"{where}: {key} is true or false, not {run[key]!r}")
    if run.get("emphasis") and run.get("bold") is False:
        raise SpecError(f"{where}: emphasis true is bold, and bold is false: say one")
    return (True if run.get("emphasis") else run.get("bold")), run.get("italic")


def _style_run(run, spec: dict, guide: Guide, *, emphasis: bool, ink_hex: str | None, fit: dict | None = None,
               bold: bool | None = None, italic: bool | None = None) -> None:
    """`fit` is the solver's resolved entry for the paragraph (layout.resolved_json's
    hints): the size and tracking it measured at, which at a scale below 1.00 are not
    the guide's. The file carries those, never the guide's scale-1 values. `bold` and
    `italic` are the run's own (run_style); None takes the role's."""
    font = run.font
    font.name = guide.face(spec)
    font.size = Pt(fit["sz"] / 100) if fit else Pt(float(spec["size_pt"]))
    font.bold = bool(spec.get("weight", 400) >= 700 or emphasis) if bold is None else bold
    font.italic = bool(spec.get("italic", False)) if italic is None else italic
    if ink_hex:
        font.color.rgb = RGBColor.from_string(ink_hex.upper())
    rpr = run._r.get_or_add_rPr()
    tracking = (fit["spc"] / 100 if fit else spec.get("tracking_pt")) or 0.0
    if tracking:
        rpr.set("spc", str(int(round(tracking * 100))))


def _paragraph_props(p, spec: dict, guide: Guide, para: dict, where: str = "text") -> None:
    p_pr = p._p.get_or_add_pPr()
    for tag in ("a:lnSpc", "a:spcBef", "a:spcAft", "a:buClr", "a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in p_pr.findall(qn(tag)):
            p_pr.remove(el)
    fit = para.get("fit")
    size = fit["sz"] / 100 if fit else float(spec["size_pt"])
    ls = spec.get("line_spacing") or 1.2  # measured default pitch (deckgen E16: 1.2002 x size, font independent)
    ln = etree.SubElement(p_pr, qn("a:lnSpc"))
    etree.SubElement(ln, qn("a:spcPts")).set("val", str(fit["lnSpc_spcPts"] if fit else int(round(ls * size * 100))))
    bef = etree.SubElement(p_pr, qn("a:spcBef"))
    # the schema minimum is 0: a negative value blanks the paragraph in LibreOffice
    etree.SubElement(bef, qn("a:spcPts")).set("val", str(max(0, fit["spcBef"] if fit else int(round((para.get("space_before_pt", spec.get("space_before_pt")) or 0.0) * 100)))))
    aft = etree.SubElement(p_pr, qn("a:spcAft"))
    etree.SubElement(aft, qn("a:spcPts")).set("val", str(max(0, fit["spcAft"] if fit else int(round((para.get("space_after_pt", spec.get("space_after_pt")) or 0.0) * 100)))))
    level = para.get("bullet_level")
    p.alignment = ALIGN.get(paragraph_align(para, spec), PP_ALIGN.LEFT)
    if level is None:
        etree.SubElement(p_pr, qn("a:buNone"))
        p_pr.set("marL", "0")
        p_pr.set("indent", "0")
    else:
        bullets = guide.g.get("typography", {}).get("bullets", {})
        entry = bullets.get(str(level)) or bullets.get(level) or bullets.get("0") or {}
        glyph = entry.get("glyph")
        if not glyph:
            glyph = "\u2022"
            guide.take(where, f"bullet level {level}: glyph \u2022 (the guide's typography.bullets name no glyph for it)", look=True)
        bufont = entry.get("font")
        # the solver measured every line at width - indent_emu (the reference's own marL for
        # the level); the file carries that marL, so the literal breaks were computed for the
        # column that is written. The hanging indent is the level's own step from the guide.
        indent = para.get("indent_emu")
        if not indent:
            role = para.get("indent_role") or ("space.sm" if "space.sm" in guide.spacing else None)
            indent = guide.space(role, 228600)
            guide.take(where, f"bullet level {level}: indent {indent} EMU ({role or 'no spacing step: 0.25 in'}; no measured indent was given)", look=role is None)
        hanging_role = entry.get("hanging_indent_role") or entry.get("indent_role")
        hanging = guide.space(hanging_role, indent) if hanging_role else indent
        # the glyph needs its own width plus a word space: a 5pt step on a 12pt body would
        # leave 0.8pt between the square and the text (reviewed on a dark PDF reference)
        hanging = max(hanging, int(size * 12700 * 0.9))
        p_pr.set("lvl", str(int(level)))
        p_pr.set("marL", str(int(indent)))
        p_pr.set("indent", str(-int(min(indent, hanging))))
        # the reference's bullet ink, where its glyphs carry their own (extract_style: bullets[level].ink)
        bu_hex = guide.hex(entry["ink_role"], where="bullet") if entry.get("ink_role") else entry.get("ink_hex")
        if bu_hex:
            clr = etree.SubElement(p_pr, qn("a:buClr"))
            etree.SubElement(clr, qn("a:srgbClr")).set("val", str(bu_hex).upper())
        if bufont:
            etree.SubElement(p_pr, qn("a:buFont")).set("typeface", bufont)
        bu = etree.SubElement(p_pr, qn("a:buChar"))
        bu.set("char", glyph)


def _body_props(txBody: etree._Element, anchor: str) -> None:
    body_pr = txBody.find(qn("a:bodyPr"))
    if body_pr is None:
        body_pr = etree.SubElement(txBody, qn("a:bodyPr"))
        txBody.insert(0, body_pr)
    for k in ("lIns", "tIns", "rIns", "bIns"):
        body_pr.set(k, "0")
    body_pr.set("wrap", "square")
    body_pr.set("anchor", {"t": "t", "m": "ctr", "b": "b"}.get(anchor, "t"))
    for tag in ("a:normAutofit", "a:spAutoFit", "a:noAutofit"):
        for el in body_pr.findall(qn(tag)):
            body_pr.remove(el)
    etree.SubElement(body_pr, qn("a:noAutofit"))


# --- element writers ---------------------------------------------------------------------------------


def write_text(slide, el: dict, guide: Guide, box: tuple[int, int, int, int], shape=None) -> None:
    x, y, w, h = box
    if shape is None:
        shape = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    else:
        # the layout's own title placeholder, moved to the solved box: the file keeps the
        # placeholder semantics (outline view, the judge's title size) and the reference's type
        shape.left, shape.top, shape.width, shape.height = Emu(x), Emu(y), Emu(w), Emu(h)
        shape.text_frame.text = ""
    shape.name = el.get("id", "text")
    if guide.text_container != "text_box" and shape.is_placeholder is False:
        # the reference draws text in plain shapes (keybank: 16 of 16 on its label page); a text box
        # is what python-pptx writes by default and what the judge counts as a different kind
        cnv = shape._element.find(f"{qn('p:nvSpPr')}/{qn('p:cNvSpPr')}")
        if cnv is not None and cnv.get("txBox") is not None:
            del cnv.attrib["txBox"]
    sp_pr = shape._element.spPr
    fill_hex = guide.hex(el.get("fill_role"), where=el.get("id", "text"))
    if fill_hex:
        _solid_fill(sp_pr, fill_hex)
    line_hex = guide.hex(el.get("line_role"), where=el.get("id", "text"))
    if el.get("line_role"):
        _set_line(sp_pr, line_hex, outline_width(el, guide))
    tf = shape.text_frame
    tf.word_wrap = True
    _body_props(tf._txBody, el.get("anchor", "t"))
    first = True
    for para in el.get("paragraphs", []):
        spec = guide.type_spec(para["role"], where=el.get("id", "text"))
        ink_hex = guide.hex(para.get("ink_role") or spec.get("ink_role"), where=para["role"]) if (para.get("ink_role") or spec.get("ink_role")) else (spec.get("ink") if HEX_RE.match(spec.get("ink") or "") else None)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        _paragraph_props(p, spec, guide, para, el.get("id", "text"))
        runs = para.get("runs") or [{"text": para.get("text", "")}]
        # the solver measured the cased text (layout.apply_case): write the same text
        runs = [{**r, "text": t} for r, t in zip(runs, apply_case_runs([r["text"] for r in runs], spec.get("case")))]
        lines = para.get("lines")
        fit = para.get("fit")
        if lines and not fit:
            raise SpecError(f"{el.get('id')}: measured lines without the solver's sizes (paragraph 'fit'); lines measured at one size and written at another overflow")
        if lines:
            # measured wrap: literal breaks between lines, runs split at line boundaries
            full = "".join(r["text"] for r in runs)
            if "".join(lines).replace(" ", "") != full.replace(" ", ""):
                raise SpecError(f"{el.get('id')}: the measured lines {lines!r} do not reconstruct the text written, {full!r}; the solver measured the text its role's case gives")
            cursor = 0
            run_spans = []
            for r in runs:
                run_spans.append((cursor, cursor + len(r["text"]), r))
                cursor += len(r["text"])
            pos = 0
            for li, line in enumerate(lines):
                start = full.find(line, pos)
                if start < 0:
                    start = pos
                end = start + len(line)
                for s0, s1, r in run_spans:
                    a, b = max(s0, start), min(s1, end)
                    if a < b:
                        pr = p.add_run()
                        pr.text = full[a:b]
                        bold, italic = run_style(r, f"{el.get('id')} run")
                        _style_run(pr, spec, guide, emphasis=False, ink_hex=_run_ink(r, guide, ink_hex), fit=fit, bold=bold, italic=italic)
                if li < len(lines) - 1:
                    p._p.append(etree.Element(qn("a:br")))
                pos = end
        else:
            for r in runs:
                pr = p.add_run()
                pr.text = r["text"]
                bold, italic = run_style(r, f"{el.get('id')} run")
                _style_run(pr, spec, guide, emphasis=False, ink_hex=_run_ink(r, guide, ink_hex), fit=fit, bold=bold, italic=italic)


#: 0.75pt: the width an outline or rule is written at when its spec and the guide state
#: none; listed under defaults_taken wherever it is used
DEFAULT_LINE_EMU = 9525


def outline_width(el: dict, guide: Guide) -> int:
    """An element's outline width: its line_width_emu, else DEFAULT_LINE_EMU, listed."""
    if el.get("line_width_emu"):
        return int(el["line_width_emu"])
    guide.take(el.get("id", "shape"), "outline width 0.75pt (line_width_emu left out)", look=True)
    return DEFAULT_LINE_EMU


def _run_ink(run: dict, guide: Guide, paragraph_ink: str | None) -> str | None:
    """A run's own ink_role over its paragraph's ink (build_deck.contrast_of scores the same)."""
    return guide.hex(run["ink_role"], where="run") if run.get("ink_role") else paragraph_ink


#: OOXML's line end types (ST_LineEndType), each a value the file states as written
LINE_ENDS = ("none", "triangle", "stealth", "diamond", "oval", "arrow")


def write_rule(slide, el: dict, guide: Guide, box: tuple[int, int, int, int]) -> None:
    """A straight line: across the middle of its box (a layout rule), or between the
    element's own `points` [x1, y1, x2, y2] (an annotation's line, leader or arrow),
    with a line end at the first point (`head_end`) and the second (`tail_end`) when they
    name one of LINE_ENDS."""
    x, y, w, h = box
    if el.get("points"):
        x1, y1, x2, y2 = (int(v) for v in el["points"])
        line = _without_theme_style(slide.shapes.add_connector(1, Emu(x1), Emu(y1), Emu(x2), Emu(y2)))
    else:
        horizontal = w >= h
        line = _without_theme_style(slide.shapes.add_connector(1, Emu(x), Emu(y + (h // 2 if horizontal else 0)), Emu(x + w if horizontal else x), Emu(y + (h // 2) if horizontal else y + h)))
    line.name = el.get("id", "rule")
    role = el.get("line_role")
    if not role:
        role = "color.rule"
        guide.take(el.get("id", "rule"), "line_role: color.rule (left out)")
    hex_color = guide.hex(role, where=el.get("id", "rule"))
    width_emu = el.get("line_width_emu")
    if not width_emu:
        # the reference's own rule in this role (a mined line motif), else the default listed
        motif = next((m for m in guide.g.get("motifs") or [] if m.get("geometry_preset") == "line" and m.get("line_role") == role and m.get("line_width_pt")), None)
        if motif is not None:
            width_emu = int(round(float(motif["line_width_pt"]) * EMU_PER_PT))
        else:
            width_emu = DEFAULT_LINE_EMU
            guide.take(el.get("id", "rule"), f"line width 0.75pt (left out, and the guide's motifs draw no {role} line)", look=True)
    _set_line(line._element.spPr, hex_color, width_emu, el.get("dash"))
    for key, tag in (("head_end", "a:headEnd"), ("tail_end", "a:tailEnd")):
        if el.get(key) not in (None, "none"):
            if el[key] not in LINE_ENDS:
                raise SpecError(f"{el.get('id', 'rule')}: {key} {el[key]!r} is not an OOXML line end ({', '.join(LINE_ENDS)})")
            etree.SubElement(line._element.spPr.find(qn("a:ln")), qn(tag)).set("type", el[key])


# A rect element names any OOXML preset geometry (prst) python-pptx writes, by its OOXML
# name. A bracket or brace is drawn opening toward what it spans: `rot` 90 turns
# leftBracket's spine to the top (a bracket above a span) and rightBracket's to the bottom
# (below it). PRESETS_DRAWN_BEFORE are the ones the emitter wrote before any preset could
# be named, drawn in the round proofs' LibreOffice renders; no preset has been measured
# against PowerPoint, and one outside this set has not been looked at in either renderer,
# so the build reports it as renderer-unmeasured (build_deck.unmeasured_presets).
PRESETS_DRAWN_BEFORE = frozenset({"rect", "roundRect", "ellipse", "chevron", "homePlate", "rightArrow", "leftArrow", "downArrow", "upArrow",
                                  "diamond", "triangle", "leftBracket", "rightBracket", "leftBrace", "rightBrace"})


def preset_shape(prst, where: str):
    """The python-pptx shape type of an OOXML preset name, or SpecError naming the miss."""
    from pptx.shapes.autoshape import AutoShapeType
    try:
        return AutoShapeType.id_from_prst(prst) if isinstance(prst, str) else None
    except ValueError:
        pass
    raise SpecError(f"{where}: prst {prst!r} is not an OOXML preset geometry python-pptx writes (named as the file names it: "
                    "rect, roundRect, ellipse, parallelogram, wedgeRectCallout, cloud, ...)")


def write_rect(slide, el: dict, guide: Guide, box: tuple[int, int, int, int]) -> None:
    x, y, w, h = box
    prst = preset_shape(el.get("prst", "rect"), el.get("id", "rect"))
    rot = int(el.get("rot") or 0) % 360
    if rot not in (0, 90, 180, 270):
        raise SpecError(f"{el.get('id', 'rect')}: rot {rot} is not a quarter turn")
    if rot in (90, 270):
        # the box is what the page shows; the shape's own frame is that box turned back
        # about its centre, so its width and height swap
        cx, cy = x + w / 2, y + h / 2
        x, y, w, h = round(cx - h / 2), round(cy - w / 2), h, w
    shape = _without_theme_style(slide.shapes.add_shape(prst, Emu(x), Emu(y), Emu(w), Emu(h)))
    if rot:
        shape.rotation = rot
    shape.name = el.get("id", "rect")
    sp_pr = shape._element.spPr
    fill_hex = guide.hex(el.get("fill_role"), where=el.get("id", "rect"))
    if fill_hex:
        _solid_fill(sp_pr, fill_hex)
    else:
        for child in list(sp_pr):
            if child.tag == qn("a:solidFill"):
                sp_pr.remove(child)
        _insert_fill_first(sp_pr, etree.Element(qn("a:noFill")))
    line_hex = guide.hex(el.get("line_role"), where=el.get("id", "rect"))
    _set_line(sp_pr, line_hex, outline_width(el, guide) if line_hex else None, el.get("dash"))
    if el.get("paragraphs"):
        tf = shape.text_frame
        tf.word_wrap = True
        _body_props(tf._txBody, el.get("anchor", "m"))
        first = True
        for para in el["paragraphs"]:
            spec = guide.type_spec(para["role"], where=el.get("id", "rect"))
            ink_hex = guide.hex(para.get("ink_role"), where=el.get("id", "rect")) if para.get("ink_role") else None
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            _paragraph_props(p, spec, guide, para, el.get("id", "rect"))
            runs = para.get("runs") or [{"text": para.get("text", "")}]
            for r, text in zip(runs, apply_case_runs([r["text"] for r in runs], spec.get("case"))):
                pr = p.add_run()
                pr.text = text
                bold, italic = run_style(r, f"{el.get('id', 'rect')} run")
                _style_run(pr, spec, guide, emphasis=False, ink_hex=_run_ink(r, guide, ink_hex), bold=bold, italic=italic)
        tf.margin_left = tf.margin_right = Emu(guide.space(el.get("pad_role"), 0))
        tf.margin_top = tf.margin_bottom = Emu(guide.space(el.get("pad_role"), 0))


#: a:blip's extension for an SVG picture (Office 2016+), with a PNG in r:embed for every
#: reader that does not draw SVG
SVG_BLIP_EXT = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
SVG_BLIP_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"


def write_image(slide, el: dict, box: tuple[int, int, int, int], guide: Guide | None = None) -> None:
    """A picture at its own aspect inside its box (images.placed_box). A file is written
    as given; an icon (svg_source) is written natively as an SVG part on the blip's svgBlip
    extension, with a PNG fallback rasterised from the same parsed shapes."""
    import images
    if "aspect" not in el:
        # a picture the build did not place through images.picture (the chrome's logo path)
        x, y, w, h = box
    else:
        x, y, w, h = images.placed_box(box, el["aspect"], el.get("align", "center"), el.get("valign", "middle"))
    if el.get("svg_source"):
        if guide is None:
            raise SpecError(f"{el.get('id')}: an icon needs the guide to resolve its colour roles")
        try:
            icon = images.parse_svg(Path(el["svg_source"]).read_bytes(), el.get("id", "image"))
            hx = images.hexes(icon, guide, el.get("id", "image"))
        except images.ImageRefused as error:
            raise SpecError(str(error)) from None
        import io
        pic = slide.shapes.add_picture(io.BytesIO(images.fallback_png(icon, hx, w, h)), Emu(x), Emu(y), Emu(w), Emu(h))
        package = slide.part.package
        svg_part = Part(package.next_partname("/ppt/media/image%d.svg"), "image/svg+xml", package, images.svg_bytes(icon, hx))
        rid = slide.part.relate_to(svg_part, RT.IMAGE)
        blip = pic._element.find(".//" + qn("a:blip"))
        ext_lst = blip.find(qn("a:extLst"))
        if ext_lst is None:
            ext_lst = etree.SubElement(blip, qn("a:extLst"))
        ext = etree.SubElement(ext_lst, qn("a:ext"))
        ext.set("uri", SVG_BLIP_EXT)
        svg_blip = etree.SubElement(ext, f"{{{SVG_BLIP_NS}}}svgBlip", nsmap={"asvg": SVG_BLIP_NS})
        svg_blip.set(qn("r:embed"), rid)
    else:
        pic = slide.shapes.add_picture(el["path"], Emu(x), Emu(y), Emu(w), Emu(h))
    pic.name = el.get("id", "image")


#: the edge classes a table's rules are written by (_horizontal_edge_class, _write_table_edges)
EDGE_CLASSES = ("table_top", "header_bottom", "totals_top", "row_interior", "table_bottom", "table_left", "column_interior", "table_right")
EDGE_RULE_KEYS = {"color_role", "width_pt", "none"}


def table_edge_rules(el: dict, guide: Guide) -> dict:
    """The table's edge rules: the guide's tables.edge_rules, with the table's own
    `edge_rules` (the guide's shape, per class: {"color_role", "width_pt"} or {"none": true})
    replacing a class it names. Roles only; a stroke under 0.5pt only at a width the guide's
    own edge rules declare (RULE_WEIGHT_FLOOR_EMU, measured)."""
    rules = dict(guide.tables.get("edge_rules") or {})
    own = el.get("edge_rules")
    if own is None:
        return rules
    where = f"{el.get('id', 'table')} edge_rules"
    if not isinstance(own, dict):
        raise SpecError(f"{where}: an object {{class: {{\"color_role\", \"width_pt\"}} or {{\"none\": true}}}}, classes {list(EDGE_CLASSES)}")
    declared = {int(r.get("width_emu") or round(float(r.get("width_pt") or 0) * EMU_PER_PT)) for r in rules.values() if isinstance(r, dict)}
    for klass, rule in own.items():
        if klass not in EDGE_CLASSES:
            raise SpecError(f"{where}: {klass!r} is not an edge class; the classes are {list(EDGE_CLASSES)}")
        if not isinstance(rule, dict) or set(rule) - EDGE_RULE_KEYS or not (rule.get("none") is True or {"color_role", "width_pt"} <= set(rule)):
            raise SpecError(f"{where}.{klass}: {{\"color_role\": <role>, \"width_pt\": <points>}} or {{\"none\": true}}, not {rule!r}")
        if rule.get("none") is not True:
            guide.hex(rule["color_role"], where=f"{where}.{klass}")
            width = rule["width_pt"]
            if not isinstance(width, (int, float)) or isinstance(width, bool) or width <= 0:
                raise SpecError(f"{where}.{klass}: width_pt is a number of points above 0")
            emu = round(float(width) * EMU_PER_PT)
            if emu < RULE_WEIGHT_FLOOR_EMU and emu not in declared:
                raise SpecError(f"{where}.{klass}: {width}pt is under 0.5pt, which PowerPoint and LibreOffice draw differently (measured), and the guide's "
                                "own table rules declare no such width")
        rules[klass] = rule
    return rules


def _table_edge(rules: dict, guide: Guide, klass: str, totals_rule: bool) -> tuple[str | None, int]:
    rule = rules.get(klass)
    if rule is None and klass == "totals_top":
        # a guide without its own totals class: the totals rule is the header rule when the
        # guide asks for one (tables.totals_rule), else the row interior like any other edge
        rule = rules.get("header_bottom") if totals_rule else rules.get("row_interior")
    if not rule or rule.get("none"):
        return None, 0
    hex_color = rule.get("color_hex") or (guide.hex(rule.get("color_role"), where=f"table edge {klass}") if rule.get("color_role") else None)
    width = rule.get("width_emu") or (int(round(float(rule["width_pt"]) * EMU_PER_PT)) if rule.get("width_pt") else 12700)
    return hex_color, int(width)


def _horizontal_edge_class(index: int, header_rows: int, n_rows: int, totals_index: int | None) -> str:
    if index == 0:
        return "table_top"
    if index == n_rows:
        return "table_bottom"
    if header_rows and index == header_rows:
        return "header_bottom"
    if totals_index is not None and index == totals_index:
        return "totals_top"
    return "row_interior"


def _write_table_edges(table, guide: Guide, n_rows: int, n_cols: int, header_rows: int, totals_index: int | None, rules: dict) -> None:
    """Every shared edge written to BOTH of its declaring cells from one rule, so the two
    sides agree byte for byte (upper lnB == lower lnT, left lnR == right lnL). Measured
    (deckgen E38 class 4, 8 of 8 cases in both axes): where the sides differ, PowerPoint
    draws the first-declared side and writes it back on save while LibreOffice draws the
    heavier one, so a totals rule declared on the totals row's lnT alone shows in the
    render and vanishes in PowerPoint. Identical sides leave no conflict to resolve."""
    def tc_pr(ri: int, ci: int) -> etree._Element:
        return table.cell(ri, ci)._tc.get_or_add_tcPr()

    for index in range(n_rows + 1):
        hex_color, width = _table_edge(rules, guide, _horizontal_edge_class(index, header_rows, n_rows, totals_index), bool(guide.tables.get("totals_rule")))
        for ci in range(n_cols):
            if index > 0:
                _cell_border(tc_pr(index - 1, ci), "a:lnB", hex_color, width)
            if index < n_rows:
                _cell_border(tc_pr(index, ci), "a:lnT", hex_color, width)
    for index in range(n_cols + 1):
        klass = "table_left" if index == 0 else ("table_right" if index == n_cols else "column_interior")
        hex_color, width = _table_edge(rules, guide, klass, bool(guide.tables.get("totals_rule")))
        for ri in range(n_rows):
            if index > 0:
                _cell_border(tc_pr(ri, index - 1), "a:lnR", hex_color, width)
            if index < n_cols:
                _cell_border(tc_pr(ri, index), "a:lnL", hex_color, width)


def _cell_border(tc_pr: etree._Element, side: str, hex_color: str | None, width_emu: int) -> None:
    for el in tc_pr.findall(qn(side)):
        tc_pr.remove(el)
    ln = etree.Element(qn(side))
    if hex_color is None or width_emu <= 0:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    else:
        ln.set("w", str(width_emu))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", hex_color.upper())
    # tcPr children order: lnL, lnR, lnT, lnB, ... then fill
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB", "a:lnTlToBr", "a:lnBlToTr"]
    idx = order.index(side)
    insert_at = 0
    for i, child in enumerate(tc_pr):
        tag = etree.QName(child).localname
        if f"a:{tag}" in order and order.index(f"a:{tag}") < idx:
            insert_at = i + 1
    tc_pr.insert(insert_at, ln)


# Table keys that set how a table is drawn, read from the table element (build_deck
# lifts them there from the tree node, its table spec or slide.styling, refusing a key
# given twice). Each fill or ink key, when present, replaces the guide's tables value
# for this table only; "none" means no fill. totals_bold and banding_phase replace the
# guide's own (table_defaults_taken lists them when neither says), edge_rules replaces
# the guide's rule for each class it names (table_edge_rules).
TABLE_STYLE_KEYS = {"totals_last_row", "emphasis", "header_align", "spanner_align",
                    "header_fill_role", "header_ink_role", "header_bold", "banding_role", "data_fill_role",
                    "totals_bold", "banding_phase", "edge_rules"}
#: One emphasis group: the body rows (data-row indices) or the columns (column keys) it
#: marks, or both (the cells where they cross: a base-case region of a sensitivity grid),
#: the fill under them (a colour role; "none", no fill, the page shows; "keep", the
#: banding or data fill the cell has without emphasis: there is no default) and optionally
#: their ink, `bold` and `italic` (true or false; left out, the cell's own: its role's
#: weight, and a totals row's totals_bold; listed under table_defaults_taken),
#: `indent_role` (a spacing step the cells' words are set in by: a sub-row) and
#: `include_header` (with columns: the header cells above them take the group's fill, ink
#: and setting too: a highlighted subject column). Emphasised cells are written and
#: measured in that setting. Several groups mark several things (the subject row shaded,
#: the row labels bold); where two groups meet, the later one's settings are drawn.
EMPHASIS_KEYS = {"rows", "columns", "fill_role", "ink_role", "bold", "italic", "indent_role", "include_header"}
#: the keys `emphasis` replaced, refused with the new spelling
OLD_EMPHASIS_KEYS = {"emphasis_rows", "emphasis_fill_role", "emphasis_ink_role"}
EMPHASIS_EXAMPLE = ('"emphasis": [{"rows": [8], "fill_role": "color.surface_alt"}, {"columns": ["label"], "fill_role": "keep", "bold": true}, '
                    '{"rows": [2, 3], "columns": ["y2025"], "fill_role": "color.surface_alt"}] (rows by data-row index, columns by key, both for the '
                    'cells where they cross; fill_role a colour role, "none" or "keep"; optional ink_role, bold, italic, indent_role, include_header)')


def table_emphasis(el: dict) -> list[dict]:
    """The table's emphasis groups, proven once: a list of objects, each naming rows or
    columns that exist, and a fill_role."""
    where = el.get("id", "table")
    old = sorted(OLD_EMPHASIS_KEYS & set(el))
    if old:
        raise SpecError(f"{where}: {old} became one key, emphasis, a list of groups: {EMPHASIS_EXAMPLE}")
    raw = el.get("emphasis")
    if raw is None:
        return []
    if not isinstance(raw, list):
        raise SpecError(f"{where}: emphasis is a list of groups: {EMPHASIS_EXAMPLE}")
    keys = [c.get("key") for c in el.get("columns") or []]
    n_rows = len(el.get("rows") or [])
    groups = []
    for i, g in enumerate(raw):
        gw = f"{where} emphasis[{i}]"
        if not isinstance(g, dict) or set(g) - EMPHASIS_KEYS or not ("rows" in g or "columns" in g):
            stray = sorted(set(g) - EMPHASIS_KEYS) if isinstance(g, dict) else []
            raise SpecError(f"{gw}: an object with rows, columns or both (both: the cells where they cross), a fill_role, and optionally "
                            f"{sorted(EMPHASIS_KEYS - {'rows', 'columns', 'fill_role'})}" + (f"; {stray} not read" if stray else "") + f": {EMPHASIS_EXAMPLE}")
        for key in ("bold", "italic", "include_header"):
            if key in g and not isinstance(g[key], bool):
                raise SpecError(f"{gw}: {key} is true or false")
        if "fill_role" not in g:
            raise SpecError(f"{gw}: needs fill_role: a colour role from the guide, \"none\" (no fill: the page shows) or \"keep\" (the banding or "
                            "data fill the cells have without emphasis); there is no default fill")
        if g.get("include_header") and "columns" not in g:
            raise SpecError(f"{gw}: include_header marks the header cells above the group's columns; name the columns")
        if "indent_role" in g and not (isinstance(g["indent_role"], str) and g["indent_role"]):
            raise SpecError(f"{gw}: indent_role is a spacing role from the guide")
        rows = cols = None
        if "rows" in g:
            rows = g["rows"]
            if not isinstance(rows, list) or not rows or any(not isinstance(r, int) or isinstance(r, bool) or not 0 <= r < n_rows for r in rows):
                raise SpecError(f"{gw}: rows {rows!r} must be data-row indices 0..{n_rows - 1}")
            rows = frozenset(rows)
        if "columns" in g:
            cols = g["columns"]
            if not isinstance(cols, list) or not cols or any(c not in keys for c in cols):
                raise SpecError(f"{gw}: columns {cols!r} must be column keys of this table ({keys})")
            cols = frozenset(keys.index(c) for c in cols)
        groups.append({"rows": rows, "columns": cols, "fill_role": g["fill_role"], "ink_role": g.get("ink_role"), "bold": g.get("bold"),
                       "italic": g.get("italic"), "indent_role": g.get("indent_role"), "include_header": bool(g.get("include_header"))})
    return groups


def emphasised(groups: list[dict], data_row: int | None, column: int) -> dict | None:
    """The group that draws a cell, the last one that marks it: a body cell (data_row its
    data-row index) by rows, columns or where they cross; a header cell (data_row None) by
    a group with include_header over its column."""
    hit = None
    for g in groups:
        if data_row is None:
            if g["include_header"] and column in g["columns"]:
                hit = g
        elif (g["rows"] is None or data_row in g["rows"]) and (g["columns"] is None or column in g["columns"]):
            hit = g
    return hit
# Keys of a table's column entries; `align` overrides the guide's
# alignment_by_column_type for that column (header_align still sets the header row).
TABLE_COLUMN_KEYS = {"key", "label", "value_kind", "width_share", "min_width_role", "align", "total"}
TABLE_ALIGNS = ("left", "center", "right")
#: what a column's totals cell is over its body (build_deck.tieouts_of checks sum and mean;
#: "none" is a figure that is no simple aggregate: a weighted average, a ratio of totals)
TABLE_TOTALS = ("sum", "mean", "none")


def table_header_bold(el: dict, guide: Guide) -> bool:
    """Whether the header row is written bold: the table's header_bold, else the guide's.
    The solver measures the header in the face this names (build_deck passes it on)."""
    return bool(el["header_bold"] if "header_bold" in el else guide.tables.get("header_bold"))


def table_roles(el: dict, guide: Guide) -> tuple[dict, dict]:
    """The header and cell type specs from the roles the solver measured, which the
    table element carries as header_role and cell_role; without them it is refused."""
    where = el.get("id", "table")
    missing = [k for k in ("header_role", "cell_role") if not el.get(k)]
    if missing:
        raise SpecError(f"{where}: table element needs {missing}, the type roles the solver measured (build_deck carries them from the layout)")
    return guide.type_spec(el["header_role"], where=f"{where} header_role"), guide.type_spec(el["cell_role"], where=f"{where} cell_role")


def table_row_styles(el: dict, guide: Guide) -> list[dict]:
    """The fill, ink and weight of every row of a table, and of every cell in it. Each
    comes from the table's own key when it has one (header_fill_role, header_ink_role,
    header_bold, banding_role, data_fill_role: a colour role or "none"), else from the
    guide's table section. A body cell an emphasis group marks (table_emphasis) takes the
    group's fill_role (a role or "none") and ink_role, else the cell ink, and is bold.
    Nothing here picks an ink for a fill: a pair that reads below 3:1 is the build's
    contrast gate's to refuse. The emitter, that gate and the annotation layer all read
    cells through this one function: each row's `cells` holds one style per column."""
    where = el.get("id", "table")
    t = guide.tables
    header_rows = int(el.get("header_rows", 1))
    n_rows = header_rows + len(el["rows"])
    for col in el["columns"]:
        extra = sorted(set(col) - TABLE_COLUMN_KEYS)
        if extra:
            raise SpecError(f"{where} column {col.get('key')!r}: unknown keys {extra}; allowed {sorted(TABLE_COLUMN_KEYS)}")
        if col.get("align") is not None and col["align"] not in TABLE_ALIGNS:
            raise SpecError(f"{where} column {col.get('key')!r}: align {col['align']!r} not in {list(TABLE_ALIGNS)}")
        if col.get("total") is not None and col["total"] not in TABLE_TOTALS:
            raise SpecError(f"{where} column {col.get('key')!r}: total {col['total']!r} not in {list(TABLE_TOTALS)}")
    for key in ("header_align", "spanner_align"):
        if el.get(key) is not None and el[key] not in TABLE_ALIGNS:
            raise SpecError(f"{where}: {key} {el[key]!r} not in {list(TABLE_ALIGNS)}")
    groups = table_emphasis(el)

    def role_hex(role: str | None, literal: str | None, label: str) -> tuple[str | None, str | None]:
        if literal:
            return literal, label
        if role and role != "none":
            return guide.hex(role, where=f"{where} {label}"), role
        return None, None

    def pick(key: str, guide_role: str | None, guide_literal: str | None, label: str) -> tuple[str | None, str | None]:
        if key in el:
            return role_hex(el[key], None, f"{label} ({key})")
        return role_hex(guide_role, guide_literal, label)

    header_fill = pick("header_fill_role", t.get("header_fill_role"), t.get("header_fill_hex"), "header fill")
    header_ink = pick("header_ink_role", t.get("header_ink_role"), None, "header ink")
    band = pick("banding_role", t.get("banding_role"), t.get("banding_hex"), "banding")
    data = pick("data_fill_role", None, t.get("data_fill_hex"), "data fill")
    _, cell_spec = table_roles(el, guide)
    cell_ink_role = cell_spec.get("ink_role")
    cell_ink = (guide.hex(cell_ink_role, where=f"{where} cell ink"), cell_ink_role) if cell_ink_role else \
        ((cell_spec["ink"], f"{el['cell_role']} observed ink") if HEX_RE.match(cell_spec.get("ink") or "") else (None, None))
    header_spec, _ = table_roles(el, guide)
    header_ink_default = header_ink if header_ink[0] else None
    styles = {id(g): (None if g["fill_role"] == "keep" else role_hex(g["fill_role"], None, f"emphasis[{k}] fill"),
                      role_hex(g["ink_role"], None, f"emphasis[{k}] ink") if g["ink_role"] else None) for k, g in enumerate(groups)}
    for g in groups:
        if g["indent_role"]:
            guide.space(g["indent_role"])  # refused by name here when the guide has no such step
    totals = bool(el.get("totals_last_row"))
    totals_bold, phase = table_totals_bold(el, guide), table_banding_phase(el, guide)
    header_bold = table_header_bold(el, guide)
    rows = []
    for ri in range(n_rows):
        is_header = ri < header_rows
        data_index = ri - header_rows
        is_totals = totals and ri == n_rows - 1
        if is_header:
            fill, ink, klass, bold, italic = header_fill, header_ink, "header", header_bold, bool(header_spec.get("italic"))
        else:
            fill = band if (band[0] and data_index % 2 == phase) else data
            ink, bold, italic = cell_ink, (totals_bold if is_totals else False) or cell_spec.get("weight", 400) >= 700, bool(cell_spec.get("italic"))
            klass = "totals" if is_totals else ("body even" if data_index % 2 == 0 else "body odd")
        base = {"fill_hex": fill[0], "fill_name": fill[1], "ink_hex": ink[0], "ink_name": ink[1], "bold": bold, "italic": italic, "indent_role": None, "class": klass}
        cells = []
        for ci in range(len(el["columns"])):
            g = emphasised(groups, None if is_header else data_index, ci)
            if g is None:
                cells.append(base)
            else:
                fill_style, ink_style = styles[id(g)]
                fh, fn = fill_style if fill_style is not None else (base["fill_hex"], base["fill_name"])
                ih, iname = ink_style if ink_style is not None else ((header_ink_default or cell_ink) if is_header else cell_ink)
                cells.append({"fill_hex": fh, "fill_name": fn, "ink_hex": ih, "ink_name": iname,
                              "bold": base["bold"] if g["bold"] is None else g["bold"], "italic": base["italic"] if g["italic"] is None else g["italic"],
                              "indent_role": g["indent_role"],
                              "class": "emphasis " + ("header " if is_header else "") + ("cells" if g["rows"] is not None and g["columns"] is not None else "rows" if g["rows"] is not None else "columns")
                                       + f" {groups.index(g)}"})
        rows.append({"row": ri, "header": is_header, **base, "cells": cells})
    return rows


def table_totals_bold(el: dict, guide: Guide) -> bool:
    """Whether a totals row is written bold: the table's totals_bold, else the guide's
    tables.totals_bold, else bold (listed under table_defaults_taken)."""
    for source in (el, guide.tables):
        if source.get("totals_bold") is not None:
            if not isinstance(source["totals_bold"], bool):
                raise SpecError(f"{el.get('id', 'table')}: totals_bold is true or false")
            return source["totals_bold"]
    return True


def table_banding_phase(el: dict, guide: Guide) -> int:
    """Which body rows the banding fill lies on: 0, the first, third, ...; 1, the second,
    fourth, ...: the table's banding_phase, else the guide's tables.banding_phase, else 0
    (listed under table_defaults_taken)."""
    for source in (el, guide.tables):
        if source.get("banding_phase") is not None:
            if source["banding_phase"] not in (0, 1) or isinstance(source["banding_phase"], bool):
                raise SpecError(f"{el.get('id', 'table')}: banding_phase is 0 (the first body row banded) or 1 (the second)")
            return int(source["banding_phase"])
    return 0


def table_defaults_taken(el: dict, guide: Guide) -> list[str]:
    """Every table choice neither the table nor the guide states, with the value written
    for it: the readers above (table_row_styles, table_totals_bold, table_banding_phase)
    and write_table's padding and alignment fall back to these. A look among them (the
    totals row's weight, which rows are banded, a column's alignment) is refused where
    nobody chose it (Guide.unchosen)."""
    out: list[str] = []
    where = f"table {el.get('id')!r}"

    def look(what: str) -> None:
        guide.unchosen(where, what)
        out.append(what)
    groups = table_emphasis(el)
    for i, g in enumerate(groups):
        if g["bold"] is None:
            out.append(f"emphasis[{i}].bold left out: the cells keep their own weight (the cell role's, and a totals row's totals_bold)")
    if el.get("totals_last_row") and el.get("totals_bold") is None and guide.tables.get("totals_bold") is None:
        look("totals_bold: the totals row is bold (neither the table nor the guide's tables says)")
    band = el.get("banding_role") if "banding_role" in el else guide.tables.get("banding_role") or guide.tables.get("banding_hex")
    if band and band != "none" and el.get("banding_phase") is None and guide.tables.get("banding_phase") is None:
        look("banding_phase: 0, the banding fill on the first body row and every second one after (neither the table nor the guide's tables says)")
    if "header_bold" not in el and guide.tables.get("header_bold") is None and int(el.get("header_rows", 1)):
        out.append("header_bold: the header row in its role's own weight (neither the table nor the guide's tables says)")
    if guide.tables.get("cell_padding_emu") is None and guide.tables.get("cell_padding_role") is None and not el.get("cell_padding_role"):
        out.append(f"cell padding: {DEFAULT_CELL_INSET_EMU} EMU (0.05 in), OOXML's own inset (the guide's tables state none)")
    align_by_kind = guide.tables.get("alignment_by_column_type") or {}
    for col in el.get("columns") or []:
        numeric = col.get("value_kind") in ("number", "percent", "currency")
        if not col.get("align") and not align_by_kind.get("numeric" if numeric else "text"):
            look(f"column {col.get('key')!r} align: {'right' if numeric else 'left'} (a {'figure' if numeric else 'text'} column; the guide's "
                       "alignment_by_column_type states none)")
    n_header = int(el.get("header_rows", 1))
    spanned = n_header > 1 and any(c1 > c0 for row in header_spans([header_labels(c, n_header, f"table {el.get('id')!r}") for c in el.get("columns") or []], n_header)
                                   for c0, c1 in row)
    if spanned and not el.get("spanner_align") and not el.get("header_align"):
        out.append("spanner_align: each spanner takes its first column's alignment (left out)")
    return out


def write_table(slide, el: dict, guide: Guide, box: tuple[int, int, int, int]) -> None:
    x, y, w, h = box
    cols = el["columns"]
    rows = el["rows"]
    header_rows = int(el.get("header_rows", 1))
    n_rows = header_rows + len(rows)
    styles = table_row_styles(el, guide)
    shape = slide.shapes.add_table(n_rows, len(cols), Emu(x), Emu(y), Emu(w), Emu(h))
    shape.name = el.get("id", "table")
    table = shape.table
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    for attr in ("firstRow", "bandRow", "lastRow", "firstCol", "lastCol", "bandCol"):
        tbl_pr.set(attr, "0")
    style_id = tbl_pr.find(qn("a:tableStyleId"))
    if style_id is not None:
        tbl_pr.remove(style_id)  # the guide decides every fill and rule; a built-in style would paint over them
    widths = el.get("column_widths_emu") or [w // len(cols)] * len(cols)
    for i, col in enumerate(table.columns):
        col.width = Emu(int(widths[i]))
    heights = el.get("row_heights_emu")
    if heights and not el.get("fit_cell"):
        raise SpecError(f"{el.get('id', 'table')}: fitted row heights without the solver's cell sizes ('fit_cell'); rows fitted at one size and written at another overflow")
    # the inset the fitter charged at the page's scale; the guide's scale-1 inset otherwise
    pad = int(el["fit_padding_emu"]) if el.get("fit_padding_emu") is not None else int(guide.tables.get("cell_padding_emu") or guide.space(guide.tables.get("cell_padding_role"), DEFAULT_CELL_INSET_EMU))
    header_spec, cell_spec = table_roles(el, guide)
    align_by_kind = guide.tables.get("alignment_by_column_type") or {}
    totals = bool(el.get("totals_last_row"))
    labels = [header_labels(c, header_rows, f"table {el.get('id')!r}") for c in cols]
    spans = header_spans(labels, header_rows)
    for ri in range(n_rows):
        row = table.rows[ri]
        if heights and ri < len(heights):
            row.height = Emu(int(heights[ri]))
        st = styles[ri]
        is_header = st["header"]
        data_index = ri - header_rows
        for ci, col in enumerate(cols):
            cell = table.cell(ri, ci)
            text = labels[ci][ri] if is_header else str(rows[data_index][ci])
            if is_header and not any(c0 == ci for c0, _ in spans[ri]):
                text = ""  # inside a spanner: the span's first cell carries its text
            cell.margin_left = cell.margin_right = cell.margin_top = cell.margin_bottom = Emu(pad)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            run = p.add_run()
            spec = header_spec if is_header else cell_spec
            run.text = apply_case(text, spec.get("case") or "as_typed")
            kind = col.get("value_kind", "label")
            align = (align_by_kind.get("numeric") if kind in ("number", "percent", "currency") else align_by_kind.get("text")) or ("right" if kind in ("number", "percent", "currency") else "left")
            if col.get("align"):
                align = col["align"]
            if is_header and el.get("header_align"):
                align = el["header_align"]
            if is_header and el.get("spanner_align") and any(c0 == ci and c1 > c0 for c0, c1 in spans[ri]):
                align = el["spanner_align"]
            p.alignment = ALIGN.get(align, PP_ALIGN.LEFT)
            cs = st["cells"][ci]
            fill_hex = cs["fill_hex"]
            # the ink is declared on the run so a dark reference does not inherit black
            _style_run(run, spec, guide, emphasis=False, ink_hex=cs["ink_hex"], fit=el.get("fit_header") if is_header else el.get("fit_cell"),
                       bold=cs["bold"], italic=cs["italic"])
            _end_paragraph_like(p, run)
            if not text:
                p._p.remove(run._r)  # an empty run is no text; the end mark alone carries the size
            ln = p._p.get_or_add_pPr()
            for tag in ("a:lnSpc", "a:spcBef", "a:spcAft"):
                for old in ln.findall(qn(tag)):
                    ln.remove(old)
            if cs["indent_role"]:
                # the solver measured these words at the cell's width less this indent
                ln.set("marL", str(int((el.get("cell_indents_emu") or {}).get(cs["indent_role"]) or guide.space(cs["indent_role"]))))
            tc_pr = cell._tc.get_or_add_tcPr()
            for child in list(tc_pr):
                if child.tag in (qn("a:solidFill"), qn("a:noFill")):
                    tc_pr.remove(child)
            if fill_hex:
                sf = etree.SubElement(tc_pr, qn("a:solidFill"))
                etree.SubElement(sf, qn("a:srgbClr")).set("val", fill_hex.upper())
            else:
                etree.SubElement(tc_pr, qn("a:noFill"))
    for ri in range(header_rows):
        for c0, c1 in spans[ri]:
            if c1 > c0:
                table.cell(ri, c0).merge(table.cell(ri, c1))  # gridSpan and hMerge: one cell across the span
    _write_table_edges(table, guide, n_rows, len(cols), header_rows, n_rows - 1 if totals else None, table_edge_rules(el, guide))


# --- charts: the spec the model writes, written as given -------------------------------------------
#
# Every design choice in a chart (type, fills, strokes, labels, axes, legend, gap and
# overlap) comes from the spec. A choice the spec leaves out is written as the neutral
# value named in CHART_DEFAULTS and listed in the note's `defaults_taken`, so the model
# reads back what it did not decide. c:varyColors val=0 is the one value written without
# being asked for: it is file safety, not design (E38 class 10, below).

CHART_KEYS = {"id", "kind", "box", "type", "categories", "series", "number_format", "text", "axes", "gap_width", "overlap", "legend", "plot_fill_role", "title", "foreign",
              "plot_area",
              "bridge"}  # bridge: a declaration the build's tie-out reads; nothing here draws from it
SERIES_KEYS = {"name", "values", "number_format", "fill_role", "line_role", "line_width_pt", "labels", "points", "marker", "smooth",
               "plot_type", "axis", "dash"}
#: the families a combo chart joins on one set of category axes (OOXML plot groups sharing
#: c:catAx); a horizontal bar, pie, scatter, bubble or radar group does not combine
COMBO_FAMILIES = ("bar", "bar_stacked", "line", "area")
SERIES_AXES = ("primary", "secondary")
LABEL_KEYS = {"show", "position", "number_format", "role", "ink_role"}
POINT_KEYS = {"index", "fill_role", "line_role", "label"}
POINT_LABEL_KEYS = LABEL_KEYS | {"text"}
AXIS_KEYS = {"visible", "line_role", "line_width_pt", "gridlines", "number_format", "reverse", "min", "max", "major_unit", "tick_marks", "labels", "cross_between", "title",
             "label_interval", "tick_interval"}
AXIS_TITLE_KEYS = {"text", "role", "ink_role", "orientation"}
#: "along": the title runs along its axis (turned on a vertical axis); "horizontal": level
AXIS_TITLE_ORIENTATIONS = ("along", "horizontal")
AXIS_NAMES = ("category", "value", "secondary_value")
# c:crossBetween on the value axis: "between" puts categories in bands (bar centres, line
# points mid-band), "midpoint" puts line points on band edges (OOXML midCat).
CROSS_BETWEEN = {"between": "between", "midpoint": "midCat"}
AXIS_LABEL_KEYS = {"show", "position", "role", "ink_role", "rotation_deg"}
TICK_LABEL_POSITIONS = {"next_to": XL_TICK_LABEL_POSITION.NEXT_TO_AXIS, "low": XL_TICK_LABEL_POSITION.LOW, "high": XL_TICK_LABEL_POSITION.HIGH}
GRIDLINE_KEYS = {"line_role", "line_width_pt"}
PLOT_AREA_KEYS = {"x", "y", "w", "h"}
DASHES = ("solid", "dot", "dash", "lgDash", "dashDot", "lgDashDot", "lgDashDotDot", "sysDash", "sysDot", "sysDashDot", "sysDashDotDot")
LEGEND_KEYS = {"position", "hide_entries", "role", "ink_role"}
TEXT_KEYS = {"role", "ink_role"}
MARKER_KEYS = {"style", "size", "fill_role", "line_role"}
TITLE_KEYS = {"text", "role", "ink_role"}
NONE = "none"
LABEL_POSITIONS = {"center": XL_LABEL_POSITION.CENTER, "inside_end": XL_LABEL_POSITION.INSIDE_END, "inside_base": XL_LABEL_POSITION.INSIDE_BASE,
                   "outside_end": XL_LABEL_POSITION.OUTSIDE_END, "above": XL_LABEL_POSITION.ABOVE, "below": XL_LABEL_POSITION.BELOW,
                   "left": XL_LABEL_POSITION.LEFT, "right": XL_LABEL_POSITION.RIGHT, "best_fit": XL_LABEL_POSITION.BEST_FIT}
LEGEND_POSITIONS = {"b": XL_LEGEND_POSITION.BOTTOM, "t": XL_LEGEND_POSITION.TOP, "r": XL_LEGEND_POSITION.RIGHT, "l": XL_LEGEND_POSITION.LEFT, "tr": XL_LEGEND_POSITION.CORNER}
#: what a guide's charts.per_kind.<kind>.legend_position takes: a chart legend position
#: (OOXML c:legendPos, as extract_style reads it) or "none"
GUIDE_LEGEND_POSITIONS = (*LEGEND_POSITIONS, NONE)


def chart_legend(el: dict):
    """The chart's legend as drawn: None (not written), "none", or its object. A legend
    object whose position is "none" (the guide's spelling) is no legend."""
    legend = el.get("legend")
    if isinstance(legend, dict) and legend.get("position") == NONE:
        return NONE
    return legend
TICK_MARKS = {"none": XL_TICK_MARK.NONE, "out": XL_TICK_MARK.OUTSIDE, "in": XL_TICK_MARK.INSIDE, "cross": XL_TICK_MARK.CROSS}
MARKER_STYLES = {m.name.lower(): m for m in XL_MARKER_STYLE if m.name not in ("PICTURE", "AUTOMATIC")}
# Label positions PowerPoint's label menu offers per family. A dLblPos outside the menu
# is a schema-valid file PowerPoint does not draw as written, so it is refused here
# rather than emitted; area, doughnut and radar labels take no position at all.
FAMILY_POSITIONS = {
    "bar": ("center", "inside_end", "inside_base", "outside_end"),
    "bar_stacked": ("center", "inside_end", "inside_base"),
    "line": ("center", "left", "right", "above", "below"),
    "xy": ("center", "left", "right", "above", "below"),
    "bubble": ("center", "left", "right", "above", "below"),
    "pie": ("center", "inside_end", "outside_end", "best_fit"),
    "doughnut": (), "area": (), "radar": (),
}
# Where a label at each position sits: on the mark's own fill, or on the plot behind it.
ON_FILL = {"bar": {"center", "inside_end", "inside_base"}, "bar_stacked": {"center", "inside_end", "inside_base"},
           "pie": {"center", "inside_end"}, "bubble": {"center"}, "area": set(), "doughnut": set()}
# The value each omitted choice is written as. "python-pptx" entries are read back from
# the chart part after it is written and reported with the value found.
CHART_DEFAULTS = {
    "series.fill_role": "no fill written: the renderer's automatic colour by series index",
    "series.line_role": "no stroke written: the renderer's automatic stroke",
    "series.line_width_pt": "no width written: the renderer's default stroke width",
    "series.labels": "no data labels",
    "labels.position": "no dLblPos: each renderer's own default position for the type",
    "labels.number_format": "linked to the source data",
    "labels.role": "the chart text (text.role, else the theme)",
    "labels.ink_role": "the chart text ink (text.ink_role, else the theme)",
    "series.marker": "the renderer's automatic marker",
    "series.smooth": "python-pptx's chart part",
    "axis.visible": "shown",
    "axis.line_role": "no stroke written: each renderer reads the theme's line style, and they differ (E38 class 5)",
    "axis.gridlines": "none (OOXML absence)",
    "axis.number_format": "linked to the source data",
    "axis.reverse": "minMax: from the origin outward (a value axis: low values at the origin; a category axis: a column chart lists its categories left to right, a horizontal bar bottom-up)",
    "axis.scale": "automatic minimum, maximum and major unit",
    "axis.tick_marks": "python-pptx's chart part",
    "axis.labels": "tick labels next to the axis in the chart text",
    "gap_width": "python-pptx's chart part",
    "overlap": "python-pptx's chart part",
    "legend": "none (OOXML absence)",
    "text": "chart text inherits the theme's face, size and colour",
    "plot_fill_role": "no plot fill: the slide shows through",
    "title": "none",
    "plot_area": "automatic: each renderer places the plot inside the frame itself",
    "axis.title": "none: no axis title",
    "axis.title.role": "no size or face written: each renderer's own axis title style (LibreOffice and PowerPoint differ)",
    "axis.title.ink_role": "no colour written: each renderer's own axis title ink",
    "axis.title.orientation": "along: the title runs along its axis, turned a quarter on a vertical axis",
    "series.axis": "primary",
}
#: the CHART_DEFAULTS that are a look: written as the theme's, a renderer's own or
#: python-pptx's, never anyone's choice (Guide.take refuses them with no reference). The
#: rest draw nothing (absence), take the chart's own text or are placed from the data.
CHART_LOOK_DEFAULTS = {"series.fill_role", "series.line_role", "series.line_width_pt", "series.marker", "series.smooth", "axis.line_role",
                       "axis.tick_marks", "gap_width", "overlap", "text", "axis.title.role", "axis.title.ink_role"}


def _keys(where: str, obj, allowed: set[str]) -> dict:
    if not isinstance(obj, dict):
        raise SpecError(f"{where}: expected an object, got {obj!r}")
    extra = sorted(set(obj) - allowed)
    if extra:
        hint = " (the chart spec changed: see references/chart-recipes.md)" if {"chart_kind", "deltas", "totals", "data_labels", "hide_value_axis"} & set(extra) else ""
        if "reference_lines" in extra:
            hint = (" (a line at a data value is a slide-level annotation now: {\"kind\": \"reference_line\", \"host\": <chart id>,"
                    " \"anchor\": {\"value\": v}}; see references/chart-recipes.md)")
        raise SpecError(f"{where}: unknown keys {extra}; allowed {sorted(allowed)}{hint}")
    return obj


def _colour(guide: Guide, role, where: str):
    """A colour role, or NONE for an explicit no-fill / no-stroke. A literal is refused."""
    if role == NONE:
        return NONE
    if not isinstance(role, str):
        raise SpecError(f"{where}: colour must be a role name or \"none\", got {role!r}")
    if HEX_RE.match(role.lstrip("#")):
        raise SpecError(f"{where}: literal colour {role!r}; a chart takes colour roles from the guide only")
    return guide.hex(role, where=where)


def unnamed_series(el: dict) -> list[int]:
    """The series a chart spec leaves without a name ("" or blank): no legend entry is drawn
    for them (LibreOffice writes "Series<n>" in its place)."""
    return [i for i, s in enumerate(el.get("series") or []) if isinstance(s, dict) and isinstance(s.get("name"), str) and not s["name"].strip()]


def marker_divergence(guide: Guide, s: dict, m: dict, where: str) -> None:
    """A marker the two renderers draw differently is refused. LibreOffice draws a marker as
    one solid colour and never its outline (measured, round 7, on an XY_SCATTER at 9pt
    markers): the series line's colour when the series draws a line, else the marker's fill,
    else the marker's line. A hollow ring (a surface fill and a coloured outline) drew
    nothing at all; PowerPoint draws it as written. So a marker is one colour: every role it
    and its series' line name is the same, and a hollow marker is an annotation (an
    outline_box with shape ellipse on the point), which both draw."""
    inks = {k: _colour(guide, r, f"{where} marker") for k, r in (("marker fill_role", m.get("fill_role")), ("marker line_role", m.get("line_role")),
                                                                ("series line_role", s.get("line_role"))) if r is not None}
    drawn = {k: v for k, v in inks.items() if v != NONE}
    hollow = inks.get("marker fill_role") == NONE and "marker line_role" in drawn
    if hollow or len(set(drawn.values())) > 1:
        named = ", ".join(f"{k} {m.get(k.split()[1]) if k.startswith('marker') else s.get('line_role')}" for k in inks)
        raise SpecError(f"{where} marker: renderer divergence: LibreOffice draws a marker as one solid colour, never its outline (the series line's "
                        f"colour when the series draws one, else the marker fill, else the marker line), so this marker ({named}) renders "
                        "differently in PowerPoint" + (" and not at all in LibreOffice" if hollow or inks.get("marker fill_role") == guide.colors.get("color.surface", {}).get("hex") else "")
                        + ". Give the marker one colour role (fill and line, and the series line when it draws one), or ring the point with an "
                        "annotation (outline_box, shape ellipse, anchor {\"point\": ...}), which both renderers draw")


def _text_style(guide: Guide, style: dict | None, where: str) -> tuple[dict | None, str | None]:
    """A text part's (type spec, ink hex) from its role and ink_role; None where omitted."""
    style = style or {}
    spec = guide.type_spec(style["role"], where=where) if style.get("role") else None
    ink = _colour(guide, style["ink_role"], where) if style.get("ink_role") else None
    if ink == NONE:
        raise SpecError(f"{where}: ink_role cannot be \"none\"")
    return spec, ink


def _font(font, guide: Guide, spec: dict | None, ink: str | None) -> None:
    if spec is not None:
        font.name = guide.face(spec)
        font.size = Pt(float(spec["size_pt"]))
        font.bold = bool(spec.get("weight", 400) >= 700)
        font.italic = bool(spec.get("italic", False))
    if ink:
        font.color.rgb = RGBColor.from_string(ink.upper())


def _fill(fmt, hex_or_none) -> None:
    if hex_or_none == NONE:
        fmt.fill.background()
    else:
        fmt.fill.solid()
        fmt.fill.fore_color.rgb = RGBColor.from_string(hex_or_none.upper())


def _stroke(fmt, hex_or_none, width_pt) -> None:
    if hex_or_none == NONE:
        fmt.line.fill.background()
        return
    fmt.line.color.rgb = RGBColor.from_string(hex_or_none.upper())
    if width_pt is not None:
        fmt.line.width = Pt(float(width_pt))


def _spelling(value, allowed) -> str:
    """A refusal's hint when the value is an allowed word in another spelling (OOXML's
    nextTo for next_to, bestFit for best_fit): case and separators ignored."""
    def norm(v) -> str:
        return re.sub(r"[^a-z0-9]", "", str(v).lower())
    same = [a for a in allowed if norm(a) == norm(value)]
    return f"; this spec spells {value!r} as {same[0]!r}" if same else ""


def _enum(where: str, label: str, value, allowed) -> None:
    if value not in allowed:
        raise SpecError(f"{where}: {label} {value!r} not in {sorted(allowed)}{_spelling(value, allowed)}")


def _check_position(family: str, position, where: str) -> None:
    if position is None:
        return
    allowed = FAMILY_POSITIONS.get(family, ())
    if position not in LABEL_POSITIONS or position not in allowed:
        raise SpecError(f"{where}: label position {position!r} is not offered for this chart type; offered: {list(allowed) or 'none (omit position)'}"
                        f"{_spelling(position, allowed or LABEL_POSITIONS)}")


def chart_parse(el: dict, guide: Guide) -> dict:
    """Validate a chart spec against the guide and list every omitted choice. Returns
    {"type": XL_CHART_TYPE member, "family", "defaults": [(path, written)]}. Both the
    emitter and the build's contrast gate read the spec through this one parse."""
    where = el.get("id", "chart")
    _keys(where, el, CHART_KEYS)
    type_name = el.get("type")
    if not isinstance(type_name, str) or type_name not in XL_CHART_TYPE.__members__:
        raise SpecError(f"{where}: type {type_name!r} is not an XL_CHART_TYPE member name (e.g. COLUMN_CLUSTERED, BAR_STACKED, LINE_MARKERS, PIE, XY_SCATTER)")
    family = chart_family(type_name)
    defaults: list[tuple[str, str]] = []

    def took(path: str, key: str) -> None:
        if key in CHART_LOOK_DEFAULTS:
            guide.take(f"{where} {path}", CHART_DEFAULTS[key], look=True)
        defaults.append((path, CHART_DEFAULTS[key]))

    series = el.get("series")
    if not isinstance(series, list) or not series:
        raise SpecError(f"{where}: series must be a non-empty list")
    cats = el.get("categories")
    if family in ("xy", "bubble"):
        if cats is not None:
            raise SpecError(f"{where}: {type_name} plots [x, y{', size' if family == 'bubble' else ''}] values; categories do not apply")
    elif not isinstance(cats, list) or not cats:
        raise SpecError(f"{where}: categories must be a non-empty list")
    width = 3 if family == "bubble" else 2
    groups = combo_groups(el, type_name, family, where)
    secondary = any(axis == "secondary" for _, axis, _ in groups)
    for si, s in enumerate(series):
        sw = f"{where} series[{si}]"
        _keys(sw, s, SERIES_KEYS)
        s_type = s.get("plot_type", type_name)
        s_family = chart_family(s_type)
        stroked = s_family in ("line", "xy", "radar")
        values = s.get("values")
        if family in ("xy", "bubble"):
            if not isinstance(values, list) or any(not isinstance(v, list) or len(v) != width for v in values):
                raise SpecError(f"{sw}: values must be [[x, y{', size' if width == 3 else ''}], ...]")
        elif not isinstance(values, list) or len(values) != len(cats):
            raise SpecError(f"{sw}: {len(values) if isinstance(values, list) else 'no'} values for {len(cats)} categories")
        if "fill_role" in s:
            if stroked:
                raise SpecError(f"{sw}: a {s_type} series draws its line_role; fill_role does not apply (marker.fill_role colours markers)")
            _colour(guide, s["fill_role"], sw)
        elif not stroked and not (family in ("pie", "doughnut") and len(s.get("points") or []) == len(values) and all("fill_role" in p for p in s.get("points") or [])):
            took(f"series[{si}].fill_role", "series.fill_role")
        if "line_role" in s:
            if _colour(guide, s["line_role"], sw) != NONE and "line_width_pt" not in s:
                took(f"series[{si}].line_width_pt", "series.line_width_pt")
        else:
            took(f"series[{si}].line_role", "series.line_role")
        if stroked and "smooth" not in s and s_family != "radar":
            took(f"series[{si}].smooth", "series.smooth")
        if "dash" in s:
            _enum(sw, "dash", s["dash"], DASHES)
            if s.get("line_role") in (None, NONE):
                raise SpecError(f"{sw}: dash styles the series' stroke; give it a line_role (and line_width_pt) to dash")
        if "marker" in s:
            m = _keys(f"{sw} marker", s["marker"], MARKER_KEYS)
            if m.get("style") is not None:
                _enum(sw, "marker style", m["style"], MARKER_STYLES)
            for k in ("fill_role", "line_role"):
                if k in m:
                    _colour(guide, m[k], f"{sw} marker")
            if m.get("style") != "none":
                marker_divergence(guide, s, m, sw)
        elif stroked and ("MARKERS" in s_type or s_type in ("XY_SCATTER", "XY_SCATTER_LINES", "XY_SCATTER_SMOOTH")):
            took(f"series[{si}].marker", "series.marker")
        labels = s.get("labels")
        if labels is None:
            took(f"series[{si}].labels", "series.labels")
        else:
            _keys(f"{sw} labels", labels, LABEL_KEYS)
            _text_style(guide, labels, f"{sw} labels")
            if labels.get("show", True):
                _check_position(s_family, labels.get("position"), f"{sw} labels")
                # a label whose point writes its own text, or hides it, shows no formatted
                # value: when every point does, the series' number format formats nothing
                own = {p.get("index") for p in s.get("points") or [] if isinstance(p, dict) and isinstance(p.get("label"), dict)
                       and ("text" in p["label"] or p["label"].get("show") is False)}
                every_own = bool(values) and all(i in own for i in range(len(values)))
                for k, d in (("position", "labels.position"), ("number_format", "labels.number_format"), ("role", "labels.role"), ("ink_role", "labels.ink_role")):
                    if k not in labels and not (k == "position" and not FAMILY_POSITIONS.get(s_family)) and not (k == "number_format" and every_own):
                        took(f"series[{si}].labels.{k}", d)
        for pi, p in enumerate(s.get("points") or []):
            pw = f"{sw} points[{pi}]"
            _keys(pw, p, POINT_KEYS)
            idx = p.get("index")
            if not isinstance(idx, int) or not 0 <= idx < len(values):
                raise SpecError(f"{pw}: index {idx!r} outside 0..{len(values) - 1}")
            for k in ("fill_role", "line_role"):
                if k in p:
                    _colour(guide, p[k], pw)
            if "label" in p:
                _keys(f"{pw} label", p["label"], POINT_LABEL_KEYS)
                _text_style(guide, p["label"], f"{pw} label")
                _check_position(s_family, p["label"].get("position"), f"{pw} label")
    if family not in ("pie", "doughnut"):
        axes = _keys(f"{where} axes", el.get("axes") or {}, set(AXIS_NAMES))
        if "secondary_value" in axes and not secondary:
            raise SpecError(f"{where} axes.secondary_value: no series is plotted on it; give a series \"axis\": \"secondary\" or remove the axis")
        cb = {n: (axes.get(n) or {}).get("cross_between") for n in ("value", "secondary_value")}
        if secondary and cb["secondary_value"] is not None and cb["secondary_value"] != cb["value"]:
            raise SpecError(f"{where} axes.secondary_value: cross_between {cb['secondary_value']!r} differs from axes.value's {cb['value']!r}; "
                            "both axes place the same categories, so the points would not line up")
        for name in ("category", "value", *(("secondary_value",) if secondary else ())):
            ax = axes.get(name)
            aw = f"{where} axes.{name}"
            if ax is None:
                took(f"axes.{name}", "axis.visible")
                ax = {}
            _keys(aw, ax, AXIS_KEYS)
            if ax.get("title") == "none":
                pass  # "none" chosen: no axis title, written as such and listed as nothing taken
            elif "title" in ax:
                if isinstance(ax["title"], str):
                    raise SpecError(f"{aw} title: {ax['title']!r}; an axis title is \"none\" (no title, chosen) or "
                                    "{\"text\": ..., \"role\": ..., \"ink_role\": ..., \"orientation\": ...}")
                t = _keys(f"{aw} title", ax["title"], AXIS_TITLE_KEYS)
                if not isinstance(t.get("text"), str) or not t["text"].strip():
                    raise SpecError(f"{aw} title: text is required")
                _text_style(guide, t, f"{aw} title")
                for k in ("role", "ink_role"):
                    if k not in t:
                        took(f"axes.{name}.title.{k}", f"axis.title.{k}")
                if "orientation" in t:
                    _enum(f"{aw} title", "orientation", t["orientation"], AXIS_TITLE_ORIENTATIONS)
                else:
                    took(f"axes.{name}.title.orientation", "axis.title.orientation")
            elif ax.get("visible") is not False:
                took(f"axes.{name}.title", "axis.title")
            if "cross_between" in ax:
                if name not in ("value", "secondary_value") or family in ("xy", "bubble"):
                    raise SpecError(f"{aw}: cross_between belongs to the value axis of a category chart" + (
                        "; an XY chart's two value axes are placed by min and max alone, and an annotation on it needs only those: remove it"
                        if family in ("xy", "bubble") else ""))
                _enum(aw, "cross_between", ax["cross_between"], CROSS_BETWEEN)
            if ax.get("visible") is False:
                # a hidden axis still scales the plot: its min, max and reverse are written
                if name == "category" and family not in ("xy", "bubble") and {"min", "max", "major_unit", "number_format"} & set(ax):
                    raise SpecError(f"{aw}: a category axis has no scale or number format")
                continue
            if "visible" not in ax and name in axes:
                took(f"axes.{name}.visible", "axis.visible")
            if "line_role" in ax:
                _colour(guide, ax["line_role"], aw)
            else:
                took(f"axes.{name}.line_role", "axis.line_role")
            if "gridlines" in ax:
                if ax["gridlines"] != NONE:
                    g = _keys(f"{aw} gridlines", ax["gridlines"], GRIDLINE_KEYS)
                    _colour(guide, g.get("line_role"), f"{aw} gridlines")
            else:
                took(f"axes.{name}.gridlines", "axis.gridlines")
            if "labels" in ax:
                _keys(f"{aw} labels", ax["labels"], AXIS_LABEL_KEYS)
                _text_style(guide, ax["labels"], f"{aw} labels")
                if ax["labels"].get("position") is not None:
                    _enum(f"{aw} labels", "position", ax["labels"]["position"], TICK_LABEL_POSITIONS)
                rot = ax["labels"].get("rotation_deg")
                if rot is not None and (isinstance(rot, bool) or not isinstance(rot, int) or not -90 <= rot <= 90):
                    raise SpecError(f"{aw} labels: rotation_deg is a whole number of degrees from -90 to 90")
            else:
                took(f"axes.{name}.labels", "axis.labels")
            if ax.get("tick_marks") is not None:
                _enum(aw, "tick_marks", ax["tick_marks"], TICK_MARKS)
            for k, d in (("tick_marks", "axis.tick_marks"), ("reverse", "axis.reverse")):
                if k not in ax:
                    took(f"axes.{name}.{k}", d)
            numeric = name in ("value", "secondary_value") or family in ("xy", "bubble")
            if numeric:
                if "number_format" not in ax:
                    took(f"axes.{name}.number_format", "axis.number_format")
                if not {"min", "max"} & set(ax):
                    took(f"axes.{name}.min/max", "axis.scale")
            elif {"min", "max", "major_unit", "number_format"} & set(ax):
                raise SpecError(f"{aw}: a category axis has no scale or number format")
            for k in ("label_interval", "tick_interval"):
                if k in ax:
                    if name != "category" or family in ("xy", "bubble"):
                        raise SpecError(f"{aw}: {k} belongs to a category axis (every n-th category labelled or ticked)")
                    if isinstance(ax[k], bool) or not isinstance(ax[k], int) or not 1 <= ax[k] <= 31999:
                        raise SpecError(f"{aw}: {k} is a whole number of categories, 1 or more")
        _values_inside_axes(el, groups, axes, family, where)
    elif el.get("axes"):
        raise SpecError(f"{where}: a {type_name} has no axes")
    if any(chart_family(t) in ("bar", "bar_stacked") for t, _, _ in groups):
        for k in ("gap_width", "overlap"):
            if k not in el:
                took(k, k)
    elif "gap_width" in el or "overlap" in el:
        raise SpecError(f"{where}: gap_width and overlap apply to bar and column types only")
    legend = el.get("legend")
    if legend is None:
        took("legend", "legend")
    elif legend != NONE:
        if not isinstance(legend, dict):
            raise SpecError(f"{where} legend: {legend!r} is neither \"none\" nor an object; write \"legend\": \"none\" for no legend, or "
                            f"\"legend\": {{\"position\": one of {sorted(LEGEND_POSITIONS)}, \"role\": ..., \"ink_role\": ...}}")
        if legend.get("position") == NONE:
            # the guide's legend_position vocabulary (GUIDE_LEGEND_POSITIONS) writes no
            # legend as position "none"; the chart takes the same word in the same place
            extra = sorted(set(legend) - {"position"})
            if extra:
                raise SpecError(f"{where} legend: position \"none\" draws no legend, so {extra} would style nothing; write \"legend\": \"none\"")
            legend = NONE
    if legend not in (None, NONE):
        _keys(f"{where} legend", legend, LEGEND_KEYS)
        _enum(f"{where} legend", "position", legend.get("position"), LEGEND_POSITIONS)
        _text_style(guide, legend, f"{where} legend")
        unnamed = unnamed_series(el)
        if unnamed and str(type_name).startswith("BAR_"):
            raise SpecError(f"{where} legend: series {unnamed} {'has' if len(unnamed) == 1 else 'have'} no name, and LibreOffice draws an unnamed "
                            "series' legend entry as 'Series<n>' (measured); on a horizontal bar chart no entry can be hidden alike in both renderers "
                            "(hide_entries below), so write \"legend\": \"none\" and name the series where they are drawn, or name every series")
        if legend.get("hide_entries") and str(type_name).startswith("BAR_"):
            # c:legendEntry/c:idx names an entry by position, and on a horizontal bar chart
            # LibreOffice 24.2 counts from the other end of the series (measured: [0] on
            # three stacked bars hid the last series, not the first): the render the agent
            # reviews would not show what the spec names
            raise SpecError(f"{where} legend: hide_entries on a horizontal bar chart ({type_name}): LibreOffice counts the entries from the other end "
                            "(measured: [0] on three stacked bars hid the last series), so the render you review hides another series than the spec names; "
                            "write \"legend\": \"none\" and name the series where they are drawn (data labels, a category_row, or text beside the chart)")
    if "text" in el:
        _keys(f"{where} text", el["text"], TEXT_KEYS)
        _text_style(guide, el["text"], f"{where} text")
    else:
        took("text", "text")
    if "plot_fill_role" in el:
        _colour(guide, el["plot_fill_role"], f"{where} plot_fill_role")
    else:
        took("plot_fill_role", "plot_fill_role")
    title = el.get("title")
    if title is None:
        took("title", "title")
    elif title != NONE:
        _keys(f"{where} title", title, TITLE_KEYS)
        _text_style(guide, title, f"{where} title")
    if "plot_area" in el:
        pa = _keys(f"{where} plot_area", el["plot_area"], PLOT_AREA_KEYS)
        if set(pa) != PLOT_AREA_KEYS or any(isinstance(v, bool) or not isinstance(v, int | float) or not 0 <= v <= 1 for v in pa.values()) \
                or pa["w"] <= 0 or pa["h"] <= 0 or pa["x"] + pa["w"] > 1 or pa["y"] + pa["h"] > 1:
            raise SpecError(f"{where} plot_area: x, y, w and h are all required, fractions of the chart frame, with x + w and y + h at most 1")
    else:
        took("plot_area", "plot_area")
    return {"type": XL_CHART_TYPE[type_name], "type_name": type_name, "family": family, "defaults": defaults, "groups": groups}


def _point_label_shown(s: dict, index: int) -> bool:
    """Whether point `index` of series `s` carries a data label: its own point label's
    show (a point label written without show is shown, as the writer draws it), else the
    series labels' show (labels absent or show false: none)."""
    for p in s.get("points") or []:
        if isinstance(p, dict) and p.get("index") == index and isinstance(p.get("label"), dict):
            return p["label"].get("show", True) is not False
    labels = s.get("labels")
    return isinstance(labels, dict) and labels.get("show", True) is not False


def _values_inside_axes(el: dict, groups: list[tuple[str, str, list[int]]], axes: dict, family: str, where: str) -> None:
    """A value past a written axis min or max is clipped at the plot's edge by the
    renderer and drawn as if it were the bound: refused, never drawn wrong. A stacked
    group is checked by its running sums (positive and negative stacks apart); a 100%
    stacked group spans its axis by construction."""
    def num(v) -> bool:
        return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)

    def check(axis_name: str, v: float, what: str) -> None:
        ax = axes.get(axis_name) or {}
        lo, hi = ax.get("min"), ax.get("max")
        if num(lo) and v < lo or num(hi) and v > hi:
            bounds = f"{lo if num(lo) else 'auto'}..{hi if num(hi) else 'auto'}"
            raise SpecError(f"{where} {what}: {v:g} is outside axes.{axis_name} {bounds}; the renderer would clip it at the plot's edge "
                            "and draw it as the bound (widen the axis, or break it and say so on the page)")

    series = el.get("series") or []
    for t, axis, members in groups:
        axis_name = "secondary_value" if axis == "secondary" else "value"
        fam = chart_family(t)
        if family in ("xy", "bubble"):
            for si in members:
                for pi, v in enumerate(series[si].get("values") or []):
                    if isinstance(v, list) and len(v) >= 2:
                        if num(v[0]):
                            check("category", v[0], f"series[{si}] point {pi} x")
                        if num(v[1]):
                            check(axis_name, v[1], f"series[{si}] point {pi} y")
            continue
        if "100" in t:
            continue
        stacked = fam == "bar_stacked" or "STACKED" in t
        n = len(el.get("categories") or [])
        for ci in range(n):
            pos = neg = 0.0
            for si in members:
                vals = series[si].get("values") or []
                v = vals[ci] if ci < len(vals) else None
                if not num(v):
                    continue
                if stacked:
                    if v >= 0:
                        pos += v
                    else:
                        neg += v
                    reach = pos if v >= 0 else neg
                    if v == 0 and not _point_label_shown(series[si], ci):
                        continue  # a zero segment draws no bar and here no label: nothing is clipped (a football field's group-header row)
                    check(axis_name, reach, f"series[{si}] category {ci} (stacked to)")
                else:
                    check(axis_name, v, f"series[{si}] category {ci}")


def combo_groups(el: dict, type_name: str, family: str, where: str) -> list[tuple[str, str, list[int]]]:
    """The chart's plot groups as (type name, value axis, series indices), the chart's own
    type on the primary axis first (possibly empty). A series names another `plot_type`
    or `axis: "secondary"` to join another group; the families combine on one category
    axis (COMBO_FAMILIES), one group per family on each axis."""
    series = el.get("series") or []
    combo = False
    for si, s in enumerate(series):
        if not isinstance(s, dict):
            continue
        if "plot_type" in s:
            if not isinstance(s["plot_type"], str) or s["plot_type"] not in XL_CHART_TYPE.__members__:
                raise SpecError(f"{where} series[{si}]: plot_type {s['plot_type']!r} is not an XL_CHART_TYPE member name")
            combo = combo or s["plot_type"] != type_name
        if "axis" in s:
            _enum(f"{where} series[{si}]", "axis", s["axis"], SERIES_AXES)
            combo = combo or s["axis"] == "secondary"
    groups: dict[tuple[str, str], list[int]] = {(type_name, "primary"): []}
    for si, s in enumerate(series):
        if isinstance(s, dict):
            groups.setdefault((s.get("plot_type", type_name), s.get("axis", "primary")), []).append(si)
    if not combo:
        return [(t, a, ix) for (t, a), ix in groups.items()]
    for t, _ in groups:
        if chart_family(t) not in COMBO_FAMILIES or t.startswith("BAR_"):
            raise SpecError(f"{where}: a combo chart joins column, line and area groups on one category axis; {t} does not combine "
                            "(a horizontal bar's categories run up the page, a pie, scatter or radar has no shared category axis)")
        if "3D" in t or t.endswith(("_PYRAMID", "_CONE", "_CYLINDER")):
            raise SpecError(f"{where}: {t} does not combine")
    for axis in SERIES_AXES:
        fams = [chart_family(t) for (t, a), ix in groups.items() if a == axis and (ix or t == type_name)]
        dup = sorted({f for f in fams if fams.count(f) > 1})
        if dup:
            raise SpecError(f"{where}: two {dup[0]} groups on the {axis} axis; on one axis each family is one group (series of one type cluster or stack together)")
    return [(t, a, ix) for (t, a), ix in groups.items()]


def _pad_point_counts(chart_space: etree._Element) -> None:
    """Every XY and bubble series declares the group's largest point count; a shorter
    series' missing points are blanks (python-pptx writes dispBlanksAs gap). Measured on
    LibreOffice 24.2: a first scatter series shorter than a later one was not drawn at
    all (10 points before 12, and 2 before 12), and drew once its ptCount matched."""
    area = chart_space.find(".//" + qn("c:plotArea"))
    for group in area if area is not None else []:
        if etree.QName(group).localname not in ("scatterChart", "bubbleChart"):
            continue
        counts = [pc for ser in group.findall(qn("c:ser")) for tag in ("c:xVal", "c:yVal", "c:bubbleSize")
                  for pc in ser.findall(f"{qn(tag)}//{qn('c:ptCount')}")]
        if counts:
            n = max(int(pc.get("val")) for pc in counts)
            for pc in counts:
                pc.set("val", str(n))


def _write_plot_layout(chart_space: etree._Element, pa: dict) -> None:
    """The plot's inner rectangle (inside the axes) as a manual layout, fractions of the
    chart frame, so a data value maps to a page position."""
    area = chart_space.find(".//" + qn("c:plotArea"))
    old = area.find(qn("c:layout"))
    if old is not None:
        area.remove(old)
    layout = etree.Element(qn("c:layout"))
    manual = etree.SubElement(layout, qn("c:manualLayout"))
    etree.SubElement(manual, qn("c:layoutTarget")).set("val", "inner")
    etree.SubElement(manual, qn("c:xMode")).set("val", "edge")
    etree.SubElement(manual, qn("c:yMode")).set("val", "edge")
    for k in ("x", "y", "w", "h"):
        etree.SubElement(manual, qn(f"c:{k}")).set("val", repr(float(pa[k])))
    area.insert(0, layout)


def _write_scale(axis, ax: dict) -> None:
    """An axis's orientation and scale as the spec gives them, shown or hidden."""
    if "reverse" in ax:
        axis.reverse_order = bool(ax["reverse"])
    if ax.get("min") is not None:
        axis.minimum_scale = float(ax["min"])
    if ax.get("max") is not None:
        axis.maximum_scale = float(ax["max"])
    if ax.get("major_unit") is not None:
        axis.major_unit = float(ax["major_unit"])
    if ax.get("cross_between"):
        _write_cross_between(axis._element, CROSS_BETWEEN[ax["cross_between"]])


def _write_cross_between(val_ax: etree._Element, value: str) -> None:
    """c:crossBetween in its schema place on a c:valAx: after c:crosses or c:crossesAt,
    before the unit and display elements. python-pptx 1.0.2 writes none on a category
    chart, so without it each renderer picks its own default."""
    old = val_ax.find(qn("c:crossBetween"))
    if old is not None:
        val_ax.remove(old)
    cb = etree.Element(qn("c:crossBetween"))
    cb.set("val", value)
    anchor = val_ax.find(qn("c:crossesAt"))
    if anchor is None:
        anchor = val_ax.find(qn("c:crosses"))
    if anchor is not None:
        anchor.addnext(cb)
    else:
        later = [c for c in val_ax if etree.QName(c).localname in ("majorUnit", "minorUnit", "dispUnits", "extLst")]
        (later[0].addprevious(cb) if later else val_ax.append(cb))


def _write_vary_colors(chart_space: etree._Element) -> None:
    """c:varyColors val=0 on every plot group that can carry it. python-pptx 1.0.2 writes
    none for a bar chart; absent, PowerPoint paints one series in five theme colours and
    LibreOffice in one (E38 class 10). With 0 both renderers draw the fills the spec
    wrote: a pie whose slices should differ names a fill per point."""
    area = chart_space.find(".//" + qn("c:plotArea"))
    if area is None:
        return
    for group in area:
        if etree.QName(group).localname not in VARY_COLORS_GROUPS:
            continue
        vc = group.find(qn("c:varyColors"))
        if vc is None:
            vc = etree.Element(qn("c:varyColors"))
            at = 0
            for i, child in enumerate(group):
                if etree.QName(child).localname in _VARY_COLORS_PREDECESSORS:
                    at = i + 1
            group.insert(at, vc)
        vc.set("val", "0")


def _write_labels(dl, guide: Guide, spec: dict, family: str, where: str) -> None:
    """One series' dLbls (python-pptx DataLabels) from a labels spec."""
    tspec, ink = _text_style(guide, spec, where)
    _font(dl.font, guide, tspec, ink)
    dl.show_value = True
    if spec.get("position") is not None:
        dl.position = LABEL_POSITIONS[spec["position"]]
    if spec.get("number_format"):
        dl.number_format = spec["number_format"]
        dl.number_format_is_linked = False


def _write_point_label(series, idx: int, guide: Guide, merged: dict, where: str) -> None:
    """A c:dLbl replaces the series' dLbls for its point, so it carries the merged
    series and point label spec in full: format, position, font, ink, text."""
    if merged.get("show") is False:
        dlbls = series._element.get_or_add_dLbls()
        d_lbl = etree.Element(qn("c:dLbl"))
        etree.SubElement(d_lbl, qn("c:idx")).set("val", str(idx))
        etree.SubElement(d_lbl, qn("c:delete")).set("val", "1")
        dlbls.insert(0, d_lbl)
        return
    dl = series.points[idx].data_label
    tspec, ink = _text_style(guide, merged, where)
    if merged.get("text") is not None:
        dl.text_frame.text = str(merged["text"])
        # every line of the label (a "\n" in the text starts a paragraph) in the label's type
        for para in dl.text_frame.paragraphs:
            for run in para.runs:
                _font(run.font, guide, tspec, ink)
    _font(dl.font, guide, tspec, ink)
    if merged.get("position") is not None:
        dl.position = LABEL_POSITIONS[merged["position"]]
    d_lbl = dl._dLbl
    if merged.get("number_format") and d_lbl is not None and merged.get("text") is None:
        fmt = etree.Element(qn("c:numFmt"))
        fmt.set("formatCode", merged["number_format"])
        fmt.set("sourceLinked", "0")
        d_lbl.find(qn("c:idx")).addnext(fmt)


def write_chart(slide, el: dict, guide: Guide, box: tuple[int, int, int, int]) -> dict:
    x, y, w, h = box
    parsed = chart_parse(el, guide)
    family = parsed["family"]
    where = el.get("id", "chart")
    if family in ("xy", "bubble"):
        data = XyChartData(number_format=el.get("number_format")) if family == "xy" else BubbleChartData(number_format=el.get("number_format"))
        for s in el["series"]:
            ser = data.add_series(s["name"], number_format=s.get("number_format"))
            for v in s["values"]:
                ser.add_data_point(*v)
    else:
        data = CategoryChartData(number_format=el.get("number_format") or "General")
        data.categories = el["categories"]
        for s in el["series"]:
            data.add_series(s["name"], s["values"], number_format=s.get("number_format"))
    try:
        gf = slide.shapes.add_chart(parsed["type"], Emu(x), Emu(y), Emu(w), Emu(h), data)
    except NotImplementedError as error:
        raise SpecError(f"{where}: python-pptx cannot write {parsed['type_name']} ({error}); choose a type it writes") from None
    gf.name = where
    chart = gf.chart
    secondary_axes = _write_combo(chart, parsed, data, where)
    if el.get("text"):
        tspec, ink = _text_style(guide, el["text"], f"{where} text")
        _font(chart.font, guide, tspec, ink)
    title = el.get("title")
    chart.has_title = bool(title and title != NONE)
    if chart.has_title:
        chart.chart_title.text_frame.text = title.get("text", "")
        tspec, ink = _text_style(guide, title, f"{where} title")
        for run in chart.chart_title.text_frame.paragraphs[0].runs:
            _font(run.font, guide, tspec, ink)
    legend = chart_legend(el)
    chart.has_legend = bool(legend and legend != NONE)
    if chart.has_legend:
        chart.legend.position = LEGEND_POSITIONS[legend["position"]]
        chart.legend.include_in_layout = False
        tspec, ink = _text_style(guide, legend, f"{where} legend")
        _font(chart.legend.font, guide, tspec, ink)
        legend_el = chart._chartSpace.find(".//" + qn("c:legend"))
        pos = legend_el.find(qn("c:legendPos"))
        # an unnamed series has no entry: LibreOffice would draw it as "Series<n>" (measured)
        for idx in sorted({int(i) for i in legend.get("hide_entries") or []} | set(unnamed_series(el)), reverse=True):
            entry = etree.Element(qn("c:legendEntry"))
            etree.SubElement(entry, qn("c:idx")).set("val", str(idx))
            etree.SubElement(entry, qn("c:delete")).set("val", "1")
            pos.addnext(entry)
    if el.get("plot_fill_role") is not None:
        plot_fmt = chart.plots[0]._element.getparent()  # c:plotArea
        sp_pr = plot_fmt.find(qn("c:spPr"))
        if sp_pr is None:
            sp_pr = etree.SubElement(plot_fmt, qn("c:spPr"))
        fill = _colour(guide, el["plot_fill_role"], f"{where} plot_fill_role")
        if fill == NONE:
            etree.SubElement(sp_pr, qn("a:noFill"))
        else:
            _solid_fill(sp_pr, fill)
    # every series by its c:idx (its place in the spec), whichever plot group holds it
    by_idx = {ser.index: ser for p in chart.plots for ser in p.series}
    written_series = [by_idx[i] for i in range(len(el["series"]))]
    for si, (series, s) in enumerate(zip(written_series, el["series"])):
        sw = f"{where} series[{si}]"
        if "fill_role" in s:
            _fill(series.format, _colour(guide, s["fill_role"], sw))
        if "line_role" in s:
            _stroke(series.format, _colour(guide, s["line_role"], sw), s.get("line_width_pt"))
        if "dash" in s:
            ln = series._element.get_or_add_spPr().find(qn("a:ln"))
            for old in ln.findall(qn("a:prstDash")):
                ln.remove(old)
            dash = etree.Element(qn("a:prstDash"))
            dash.set("val", s["dash"])
            # a:prstDash follows the line's fill in CT_LineProperties
            fill = next((c for c in ln if etree.QName(c).localname in ("noFill", "solidFill", "gradFill", "pattFill")), None)
            (fill.addnext(dash) if fill is not None else ln.insert(0, dash))
        if "smooth" in s:
            # through the element: python-pptx's XySeries has no smooth property
            series._element.get_or_add_smooth().val = bool(s["smooth"])
        if "marker" in s:
            m = s["marker"]
            if m.get("style") is not None:
                series.marker.style = MARKER_STYLES[m["style"]]
            if m.get("size") is not None:
                series.marker.size = int(m["size"])
            if "fill_role" in m:
                _fill(series.marker.format, _colour(guide, m["fill_role"], f"{sw} marker"))
            if "line_role" in m:
                _stroke(series.marker.format, _colour(guide, m["line_role"], f"{sw} marker"), None)
        labels = s.get("labels")
        if labels and labels.get("show", True):
            _write_labels(DataLabels(series._element.get_or_add_dLbls()), guide, labels, family, f"{sw} labels")
        for p in s.get("points") or []:
            pw = f"{sw} point {p['index']}"
            if "fill_role" in p:
                _fill(series.points[p["index"]].format, _colour(guide, p["fill_role"], pw))
            if "line_role" in p:
                _stroke(series.points[p["index"]].format, _colour(guide, p["line_role"], pw), s.get("line_width_pt"))
            if "label" in p:
                merged = {**(labels or {}), **p["label"]}
                if not (labels and labels.get("show", True)) and "show" not in p["label"]:
                    merged["show"] = True
                _write_point_label(series, p["index"], guide, merged, f"{pw} label")
    taken = dict(parsed["defaults"])
    for si, (series, s) in enumerate(zip(written_series, el["series"])):
        if f"series[{si}].smooth" in taken:
            smooth = series._element.smooth
            written = "absent" if smooth is None else str(bool(smooth.val))
            taken[f"series[{si}].smooth"] = f"{written} ({CHART_DEFAULTS['series.smooth']})"
    # every bar group takes the chart's gap_width and overlap (a combo's columns on either axis)
    for plot in [p for p in chart.plots if etree.QName(p._element).localname == "barChart"]:
        if "gap_width" in el:
            plot.gap_width = int(el["gap_width"])
        else:
            taken["gap_width"] = f"{plot.gap_width} ({CHART_DEFAULTS['gap_width']})"
        if "overlap" in el:
            plot.overlap = int(el["overlap"])
        else:
            taken["overlap"] = f"{plot.overlap} ({CHART_DEFAULTS['overlap']})"
    if family not in ("pie", "doughnut"):
        axes = el.get("axes") or {}
        named = [("category", chart.category_axis), ("value", chart.value_axis)]
        if secondary_axes is not None:
            # python-pptx's value_axis is the SECOND c:valAx whenever there are two (its XY
            # reading); a combo's primary value axis is the first
            from pptx.chart.axis import ValueAxis
            named = [("category", chart.category_axis), ("value", ValueAxis(chart._chartSpace.valAx_lst[0])), ("secondary_value", secondary_axes)]
        for name, axis in named:
            ax = axes.get(name) or {}
            if name == "secondary_value" and "cross_between" not in ax and (axes.get("value") or {}).get("cross_between"):
                # the same categories on both pairs of axes: the secondary places them as the primary does
                ax = {**ax, "cross_between": axes["value"]["cross_between"]}
            aw = f"{where} axes.{name}"
            if ax.get("visible") is False and name == "secondary_value":
                # LibreOffice drops the min and max of a DELETED secondary axis and plots its
                # series on the primary scale (measured, round 7: a 0..30 margin line drawn
                # flat along a 0..250 baseline); a deleted primary axis keeps its scale. So
                # the secondary is written present and drawn as nothing: no line, no tick
                # marks, no labels
                axis.has_major_gridlines = False
                axis.has_minor_gridlines = False
                _stroke(axis.format, NONE, None)
                axis.major_tick_mark = TICK_MARKS["none"]
                axis.minor_tick_mark = TICK_MARKS["none"]
                axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
                _write_scale(axis, ax)
                continue
            if ax.get("visible") is False:
                axis.visible = False
                axis.has_major_gridlines = False
                # deleted, the axis still sets the plot's scale; without its min and max the
                # renderer autoscales and every computed data position is wrong
                _write_scale(axis, ax)
                continue
            # python-pptx's chart part writes value-axis gridlines; an omitted choice is none
            axis.has_major_gridlines = False
            axis.has_minor_gridlines = False
            if "line_role" in ax:
                _stroke(axis.format, _colour(guide, ax["line_role"], aw), ax.get("line_width_pt"))
            g = ax.get("gridlines")
            if g and g != NONE:
                axis.has_major_gridlines = True
                _stroke(axis.major_gridlines.format, _colour(guide, g["line_role"], f"{aw} gridlines"), g.get("line_width_pt"))
            lab = ax.get("labels")
            if lab is not None:
                if lab.get("show", True) is False:
                    axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
                else:
                    tspec, ink = _text_style(guide, lab, f"{aw} labels")
                    _font(axis.tick_labels.font, guide, tspec, ink)
                    if lab.get("position") is not None:
                        axis.tick_label_position = TICK_LABEL_POSITIONS[lab["position"]]
            if "tick_marks" in ax:
                axis.major_tick_mark = TICK_MARKS[ax["tick_marks"]]
            else:
                taken[f"axes.{name}.tick_marks"] = f"{axis.major_tick_mark.name.lower()} ({CHART_DEFAULTS['axis.tick_marks']})"
            if ax.get("number_format"):
                axis.tick_labels.number_format = ax["number_format"]
                axis.tick_labels.number_format_is_linked = False
            if isinstance(ax.get("title"), dict):
                _write_axis_title(axis, ax["title"], guide, name, chart, f"{aw} title")
            if (ax.get("labels") or {}).get("rotation_deg") is not None:
                body = axis._element.get_or_add_txPr().find(qn("a:bodyPr"))
                body.set("rot", str(int(ax["labels"]["rotation_deg"]) * 60000))
                body.set("vert", "horz")
            for k, tag in (("label_interval", "c:tickLblSkip"), ("tick_interval", "c:tickMarkSkip")):
                if k in ax:
                    _write_skip(axis._element, tag, int(ax[k]))
            _write_scale(axis, ax)
    _write_vary_colors(chart._chartSpace)
    _pad_point_counts(chart._chartSpace)
    if el.get("plot_area"):
        _write_plot_layout(chart._chartSpace, el["plot_area"])
    return {"id": where, "type": parsed["type_name"], "defaults_taken": [f"{k}: {v}" for k, v in taken.items()]}


#: c:axId values for a combo's secondary axes: any unsigned int not already used; python-pptx
#: numbers its own axes with 8-digit ids, so these cannot collide
SECONDARY_AX_IDS = ("50010", "50020")


def _write_combo(chart, parsed: dict, data, where: str):
    """Move each series that names another plot_type or the secondary axis into its own
    plot group. The group is taken from a chart python-pptx writes of that type over the
    SAME chart data (so every series' cell references are the embedded workbook's), keeps
    only its own series, and is joined to the primary axes or to a pair of secondary axes
    (a deleted category axis and a value axis on the right that crosses at the maximum).
    Returns the secondary value axis (python-pptx ValueAxis) or None."""
    groups = parsed["groups"]
    if len(groups) == 1 and groups[0][1] == "primary":
        return None
    from pptx.chart.axis import ValueAxis
    area = chart._chartSpace.find(".//" + qn("c:plotArea"))
    main = chart.plots[0]._element
    main_ids = [a.get("val") for a in main.findall(qn("c:axId"))]
    axes_first = next(c for c in area if etree.QName(c).localname in ("catAx", "valAx", "dateAx", "serAx"))
    secondary_val = None
    sec_needed = any(a == "secondary" for _, a, ix in groups if ix)
    if sec_needed:
        cat = area.find(qn("c:catAx"))
        val = area.find(qn("c:valAx"))
        s_cat, s_val = copy.deepcopy(cat), copy.deepcopy(val)
        s_cat.find(qn("c:axId")).set("val", SECONDARY_AX_IDS[0])
        s_cat.find(qn("c:crossAx")).set("val", SECONDARY_AX_IDS[1])
        delete = s_cat.find(qn("c:delete"))
        delete.set("val", "1")
        s_val.find(qn("c:axId")).set("val", SECONDARY_AX_IDS[1])
        s_val.find(qn("c:crossAx")).set("val", SECONDARY_AX_IDS[0])
        s_val.find(qn("c:axPos")).set("val", "r")
        for g in s_val.findall(qn("c:majorGridlines")) + s_val.findall(qn("c:minorGridlines")):
            s_val.remove(g)
        crosses = s_val.find(qn("c:crosses"))
        if crosses is None:
            raise SpecError(f"{where}: the value axis python-pptx wrote has no c:crosses to set for the secondary axis")
        crosses.set("val", "max")
        last_axis = [c for c in area if etree.QName(c).localname in ("catAx", "valAx", "dateAx", "serAx")][-1]
        last_axis.addnext(s_cat)
        s_cat.addnext(s_val)
        secondary_val = ValueAxis(s_val)
    for type_name, axis, idxs in groups:
        if (type_name, axis) == (parsed["type_name"], "primary"):
            continue
        if not idxs:
            continue
        from pptx import Presentation as _P
        scratch = _P()
        sslide = scratch.slides.add_slide(scratch.slide_layouts[6])
        other = sslide.shapes.add_chart(XL_CHART_TYPE[type_name], Emu(0), Emu(0), Emu(914400), Emu(914400), data).chart
        group = copy.deepcopy(other.plots[0]._element)
        for ser in group.findall(qn("c:ser")):
            if int(ser.find(qn("c:idx")).get("val")) not in idxs:
                group.remove(ser)
        ids = main_ids if axis == "primary" else list(SECONDARY_AX_IDS)
        for a, v in zip(group.findall(qn("c:axId")), ids):
            a.set("val", v)
        axes_first.addprevious(group)
    for ser in main.findall(qn("c:ser")):
        if int(ser.find(qn("c:idx")).get("val")) not in groups[0][2]:
            main.remove(ser)
    if not main.findall(qn("c:ser")):
        # every series left the chart's own type: its group would plot nothing
        area.remove(main)
    # groups in the order they are painted: areas under bars under lines. LibreOffice paints
    # in document order (measured by the probe: markers of a line written before a bar group
    # were hidden under the bars), so the file states the order instead of leaving it to it
    rank = {"areaChart": 0, "barChart": 1, "lineChart": 2}
    groups_el = [c for c in area if etree.QName(c).localname in rank]
    anchor = groups_el[-1].getnext()
    for g in groups_el:
        area.remove(g)
    for g in sorted(groups_el, key=lambda g: rank[etree.QName(g).localname]):
        anchor.addprevious(g)
    return secondary_val


def _write_skip(cat_ax: etree._Element, tag: str, n: int) -> None:
    """c:tickLblSkip or c:tickMarkSkip in its schema place on a c:catAx: after lblOffset
    (tickLblSkip before tickMarkSkip), before noMultiLvlLbl and extLst."""
    for old in cat_ax.findall(qn(tag)):
        cat_ax.remove(old)
    el = etree.Element(qn(tag))
    el.set("val", str(n))
    later = ("tickMarkSkip", "noMultiLvlLbl", "extLst") if tag == "c:tickLblSkip" else ("noMultiLvlLbl", "extLst")
    nxt = next((c for c in cat_ax if etree.QName(c).localname in later), None)
    (nxt.addprevious(el) if nxt is not None else cat_ax.append(el))


def _write_axis_title(axis, title: dict, guide: Guide, name: str, chart, where: str) -> None:
    axis.has_title = True
    tf = axis.axis_title.text_frame
    tf.text = title["text"]
    tspec, ink = _text_style(guide, title, where)
    for run in tf.paragraphs[0].runs:
        _font(run.font, guide, tspec, ink)
    vertical_axis = name in ("value", "secondary_value") and not str(chart.chart_type).startswith("BAR_")
    orientation = title.get("orientation", "along")
    body = axis.axis_title._title.find(".//" + qn("c:rich")).find(qn("a:bodyPr"))
    # written either way: the renderers' absence values for a vertical axis title differ
    body.set("rot", "-5400000" if vertical_axis and orientation == "along" else "0")
    body.set("vert", "horz")


def chart_label_inks(el: dict, guide: Guide) -> list[dict]:
    """Every chart text the spec inks, with what it sits on, for the build's contrast
    gate: data labels per series and per overriding point, axis tick labels, legend.
    `on` is a list of ("fill", hex, name) and ("plot",); "plot" is the plot fill when
    the spec names one, else whatever the page draws under the chart. A label whose
    fill the spec left to the renderer is ("unknown", why). A position the spec left
    open is scored against every backing the label can land on."""
    parsed = chart_parse(el, guide)
    family = parsed["family"]
    where = el.get("id", "chart")
    text_spec, text_ink = _text_style(guide, el.get("text"), f"{where} text")
    def fam(si: int) -> str:
        return chart_family(el["series"][si].get("plot_type", parsed["type_name"]))
    out: list[dict] = []

    def ink_of(spec: dict) -> tuple[str | None, str | None, dict | None]:
        tspec, ink = _text_style(guide, spec, where)
        name = spec.get("ink_role") or ((el.get("text") or {}).get("ink_role"))
        return ink or text_ink, name, tspec or text_spec

    def backing(fill_role, position, family: str) -> list[tuple]:
        on_fill = ON_FILL.get(family, set())
        if family in ("line", "xy", "radar"):
            return [("plot",)]
        if fill_role is None:
            fill = ("unknown", "the series fill is left to the renderer; name a fill_role to score it")
        elif fill_role == NONE:
            fill = ("plot",)
        else:
            fill = ("fill", guide.hex(fill_role, where=where), fill_role)
        if family in ("area", "doughnut"):
            return [fill]
        if position is None:
            return [fill, ("plot",)]  # the renderer picks the position
        return [fill] if position in on_fill else [("plot",)]

    for si, s in enumerate(el["series"]):
        labels = s.get("labels")
        series_on = labels is not None and labels.get("show", True)
        overridden = {}
        for p in s.get("points") or []:
            if "label" in p or "fill_role" in p:
                overridden[p["index"]] = p
        n = len(s["values"])
        plain = [i for i in range(n) if i not in overridden]
        if series_on and plain:
            ink, name, tspec = ink_of(labels)
            out.append({"element": where, "series": si, "part": f"series[{si}] {s['name']!r} labels", "points": plain, "ink": ink, "ink_name": name,
                        "spec": tspec, "on": backing(s.get("fill_role"), labels.get("position"), fam(si))})
        for idx, p in sorted(overridden.items()):
            merged = {**(labels or {}), **p.get("label", {})}
            shown = merged.get("show", True) if ("label" in p or series_on) else False
            if not shown:
                continue
            ink, name, tspec = ink_of(merged)
            out.append({"element": where, "series": si, "part": f"series[{si}] {s['name']!r} point {idx} label", "points": [idx], "ink": ink, "ink_name": name,
                        "spec": tspec, "on": backing(p.get("fill_role", s.get("fill_role")), merged.get("position"), fam(si))})
    if family not in ("pie", "doughnut"):
        for name in ("category", "value", *(("secondary_value",) if any(x.get("axis") == "secondary" for x in el["series"]) else ())):
            ax = (el.get("axes") or {}).get(name) or {}
            if ax.get("visible") is False or (ax.get("labels") or {}).get("show") is False:
                continue
            ink, ink_name, tspec = ink_of(ax.get("labels") or {})
            out.append({"element": where, "part": f"{name} axis labels", "points": [], "ink": ink, "ink_name": ink_name, "spec": tspec, "on": [("plot",)]})
    legend = chart_legend(el)
    if legend and legend != NONE:
        ink, ink_name, tspec = ink_of(legend)
        out.append({"element": where, "part": "legend", "points": [], "ink": ink, "ink_name": ink_name, "spec": tspec, "on": [("plot",)]})
    return out


# --- slides ------------------------------------------------------------------------------------------


def _layout_named(prs, name: str | None):
    layouts = [lay for m in prs.slide_masters for lay in m.slide_layouts]
    if name:
        for lay in layouts:
            if lay.name.casefold() == name.casefold():
                return lay
        raise SpecError(f"layout {name!r} not in template; available: {[lay.name for lay in layouts]}")
    for preferred in ("Title and Content", "Title Only", "Blank"):
        for lay in layouts:
            if lay.name.casefold() == preferred.casefold():
                return lay
    return min(layouts, key=lambda lay: len(lay.placeholders))


def _copy_chrome_placeholders(slide, layout, wanted: set[str]) -> None:
    """python-pptx clones only title/body placeholders; date, footer and slide-number
    placeholders must be copied by hand or the page shows no number."""
    for sp in layout.placeholders:
        kind = str(sp.placeholder_format.type).split(".")[-1].split(" ")[0]
        if kind in wanted:
            el = copy.deepcopy(sp._element)
            slide.shapes._spTree.append(el)
            # keep the field text but re-id the shape
            cnv = el.find(".//" + qn("p:cNvPr"))
            if cnv is not None:
                cnv.set("id", str(max(int(c.get("id", "0")) for c in slide.shapes._spTree.iter(qn("p:cNvPr"))) + 1))


def picture_file(path, where: str) -> str:
    """A picture file named by a guide or a deck (a chrome logo), read through the one
    definition of a picture (images.read_picture): its format extension, or SpecError."""
    import images
    if not isinstance(path, str) or not path:
        raise SpecError(f"{where}: a file path")
    try:
        return images.read_picture(path, where)[0]
    except images.ImageRefused as error:
        raise SpecError(str(error)) from None


def chrome_hex(ch: dict, guide: Guide, part: str) -> str | None:
    """An entry's observed `<part>_hex`, else the hex of its `<part>_role` (part: fill or
    line). None when the entry names neither: nothing is chosen for it."""
    if ch.get(f"{part}_hex"):
        return ch[f"{part}_hex"]
    role = ch.get(f"{part}_role")
    return guide.colors[role]["hex"] if role in guide.colors else None


CHROME_TEXT_KINDS = ("label", "page_number", "footer", "date")


def chrome_text(ch: dict) -> str | None:
    """The words a text entry draws, without its page-number marker; None for an entry
    that draws no text."""
    if ch["kind"] not in CHROME_TEXT_KINDS or not (ch.get("text") or ch["kind"] == "page_number"):
        return None
    return (ch.get("text") or "").replace("<page number>", "").strip()


def recurring_chrome(guide_data: dict) -> list[dict]:
    """The guide's chrome entries a generated page can carry, the one definition: every entry
    on at least half the reference's slides; on a guide made with no reference, every entry
    (each the model's choice for every page)."""
    import plan
    if not plan.has_reference(guide_data):
        return list(guide_data.get("chrome") or [])
    n = max(1, guide_data.get("reference", {}).get("slides", 1))
    return [ch for ch in guide_data.get("chrome", []) if ch.get("slides_present", 0) >= 0.5 * n]


def guide_chrome(guide_data: dict, mode: str) -> list[dict]:
    """The guide's chrome a generated page starts from, in the reference's paint order:
    recurring_chrome, and under --chrome none only what the template inherits (an entry on a
    layout or master). An entry without `paint_order` (a hand-written guide) keeps its list
    place."""
    entries = recurring_chrome(guide_data)
    if mode == "none":
        entries = [ch for ch in entries if _layers(ch) != ["slide"]]
    return [ch for _, ch in sorted(enumerate(entries), key=lambda ie: (ie[1].get("paint_order", ie[0]), ie[0]))]


#: words that make a chrome line a legal or confidentiality mark: never left off a page
LEGAL_MARK = re.compile(r"confidential|proprietary|privileged|copyright|\u00a9|\(c\)|all rights reserved|strictly|not for (?:distribution|circulation)|"
                        r"do not (?:distribute|circulate|copy)|draft|disclaimer|for discussion purposes", re.IGNORECASE)
PAGE_CHROME_KEYS = {"omit", "text", "add", "why"}


def minority_chrome(guide_data: dict) -> dict[str, dict]:
    """The reference's chrome entries on fewer than half its pages (recurring_chrome leaves them
    off): a page may add one, as the reference's own minority pages do."""
    drawn = {ch["chrome_id"] for ch in recurring_chrome(guide_data)}
    return {ch["chrome_id"]: ch for ch in guide_data.get("chrome") or [] if ch["chrome_id"] not in drawn}


def added_chrome(spec: dict) -> set[str]:
    """The minority entries some page of the deck adds (its chrome {id: {"add": true}})."""
    return {cid for sd in spec.get("slides") or [] if isinstance(sd.get("chrome"), dict)
            for cid, choice in sd["chrome"].items() if isinstance(choice, dict) and choice.get("add") is True}


def page_chrome(guide_data: dict, mode: str, raw, where: str) -> dict[str, dict]:
    """A page's own chrome choices, proven at the boundary: {chrome_id: {"omit": true,
    "why": words}} leaves a reference entry off this page, {chrome_id: {"text": words, "why":
    words}} sets another line in a text entry's place, {chrome_id: {"add": true, "why": words}}
    draws on this page an entry the reference carries on fewer than half its pages
    (minority_chrome; the deck's firm_chrome choice applies to it as to any other). Only an
    entry the page would draw (guide_chrome) can be left off or reworded; never the firm's chrome (decided for the whole deck in
    firm_chrome) or a legal or confidentiality mark; under --chrome auto, never an entry the
    page inherits from its layout or master (a page cannot leave out what its layout draws:
    use a layout without it)."""
    if raw is None:
        return {}
    given = {ch["chrome_id"]: ch for ch in guide_chrome(guide_data, mode)}
    minority = minority_chrome(guide_data)
    listing = ", ".join(f"{cid} ({ch['kind']}{': ' + repr(chrome_text(ch)) if chrome_text(ch) else ''}, "
                        + (f"on {ch.get('slides_present')} of {(guide_data.get('reference') or {}).get('slides')} reference pages)"
                           if ch.get("slides_present") is not None else "chosen for every page)") for cid, ch in given.items()) or "none"
    if minority:
        listing += "; on fewer than half its pages, which a page may add: " + ", ".join(
            f"{cid} ({ch['kind']}{': ' + repr(chrome_text(ch)) if chrome_text(ch) else ''}, on {ch.get('slides_present')} pages)" for cid, ch in minority.items())
    if not isinstance(raw, dict):
        raise SpecError(f"{where} chrome: an object {{chrome_id: {{\"omit\": true, \"why\": ...}} | {{\"text\": ..., \"why\": ...}}}}; the reference's chrome here: {listing}")
    firm = {ch["chrome_id"] for ch in firm_chrome(guide_data, mode)}
    out = {}
    for cid, choice in raw.items():
        ch = given.get(cid) or minority.get(cid)
        at = f"{where} chrome {cid}"
        if ch is None:
            raise SpecError(f"{at}: not chrome of this guide; the reference's chrome here: {listing}")
        if not isinstance(choice, dict) or set(choice) - PAGE_CHROME_KEYS or sum(k in choice for k in ("omit", "text", "add")) != 1:
            raise SpecError(f"{at}: {{\"omit\": true, \"why\": ...}}, {{\"text\": ..., \"why\": ...}} or {{\"add\": true, \"why\": ...}}, not {choice!r}")
        why = choice.get("why")
        if not (isinstance(why, str) and len(why.split()) >= 3):
            raise SpecError(f"{at}: why, in words: what this page does that the reference's {ch.get('slides_present')} pages with it do not")
        if "add" in choice:
            if choice["add"] is not True:
                raise SpecError(f"{at}: add is true")
            if cid in given:
                raise SpecError(f"{at}: already drawn on every page; add takes an entry on fewer than half the reference's pages: {listing}")
            if _layers(ch) != ["slide"] and mode != "draw":
                raise SpecError(f"{at}: the reference draws it on its {'/'.join(_layers(ch))}; choose that layout, or build with --chrome draw")
            out[cid] = {"add": True, "why": why}
            continue
        if cid not in given:
            raise SpecError(f"{at}: not chrome this page draws (on {ch.get('slides_present')} reference pages); add it with {{\"add\": true, \"why\": ...}}")
        if cid in firm:
            raise SpecError(f"{at}: the firm's own chrome is kept, dropped or replaced for the whole deck in firm_chrome, never page by page")
        if LEGAL_MARK.search(chrome_text(ch) or ""):
            raise SpecError(f"{at}: {chrome_text(ch)!r} is a legal or confidentiality mark; it stays on every page that carries the reference's chrome")
        if mode == "auto" and _layers(ch) != ["slide"]:
            raise SpecError(f"{at}: the page inherits it from its {'/'.join(_layers(ch))}; choose a layout without it")
        if "omit" in choice:
            if choice["omit"] is not True:
                raise SpecError(f"{at}: omit is true")
            out[cid] = {"omit": True, "why": why}
        else:
            if chrome_text(ch) is None or not isinstance(choice["text"], str) or not choice["text"].strip():
                raise SpecError(f"{at}: text replaces the words of a text entry (label, footer, date), in words")
            if LEGAL_MARK.search(choice["text"]):
                raise SpecError(f"{at}: a legal or confidentiality mark is the reference's, not a page's own line")
            out[cid] = {"text": choice["text"].strip(), "why": why}
    return out


def chrome_on_page(guide_data: dict, mode: str, page: dict | None = None) -> list[dict]:
    """The chrome a generated page carries: guide_chrome as the page's own choices
    (page_chrome, already proven) leave it, an omitted entry off and a replaced one with
    its new words. The emitter draws from this list and build_deck reserves zones and
    scores contrast from it, so what is checked is what is drawn: an omitted footer frees
    its zone, a replaced line is scored in its new words."""
    out = []
    for ch in guide_chrome(guide_data, mode):
        choice = (page or {}).get(ch["chrome_id"])
        if choice is None:
            out.append(ch)
        elif "text" in choice:
            out.append({**ch, "text": choice["text"] + (" <page number>" if "<page number>" in (ch.get("text") or "") else "")})
    added = [ch for ch in minority_chrome(guide_data).values() if ((page or {}).get(ch["chrome_id"]) or {}).get("add")]
    if added:
        order = {ch["chrome_id"]: i for i, ch in enumerate(guide_data.get("chrome") or [])}
        out = sorted(out + added, key=lambda ch: (ch.get("paint_order", order[ch["chrome_id"]]), order[ch["chrome_id"]]))
    return out


def draw_chrome(slide, guide: Guide, guide_data: dict, page_w: int, page_h: int, entries: list[dict] | None = None, notes: list[dict] | None = None) -> list[str]:
    """The guide's own chrome, drawn from the observed hexes and positions. These are
    observed values of the reference, which is why literal hexes are allowed here and
    nowhere else. `entries` is the list to draw, in paint order; by default
    chrome_on_page(guide_data, "draw") (the PDF and blank-template path, --chrome draw).
    An entry is drawn at most once per slide and named by its chrome_id. A text entry is
    one line at its observed face, size and ink, turned by its rotation_deg. A picture
    entry (kind logo) becomes a picture when the entry carries `image_path`, else a rect
    placeholder in the entry's own fill or line, and `notes` names it; an entry with
    neither is not drawn, and `notes` says so. Returns the chrome_ids drawn."""
    if entries is None:
        entries = chrome_on_page(guide_data, "draw")
    present = {sp.name for sp in slide.shapes}
    drawn: list[str] = []
    for ch in entries:
        if ch["chrome_id"] in present:
            continue
        x, y, w, h = (int(v * 914400) for v in ch["box_in"])
        kind = ch["kind"]
        text = chrome_text(ch)
        if kind == "rule":
            if ch.get("line_width_pt"):
                width_emu = int(round(float(ch["line_width_pt"]) * EMU_PER_PT))
            else:
                width_emu = DEFAULT_LINE_EMU
                guide.take(ch["chrome_id"], "chrome rule width 0.75pt (the entry records no line_width_pt)", look=True)
            line = _without_theme_style(slide.shapes.add_connector(1, Emu(x), Emu(y + h // 2), Emu(x + w), Emu(y + h // 2)))
            line.name = ch["chrome_id"]
            ink = chrome_hex(ch, guide, "line") or chrome_hex(ch, guide, "fill")
            if not ink:
                ink = guide.colors.get("color.rule", {}).get("hex") or "000000"
                guide.take(ch["chrome_id"], f"chrome rule colour {ink} ({'color.rule' if guide.colors.get('color.rule') else 'black: the guide has no color.rule'}; the entry records no line or fill)",
                           look=not guide.colors.get("color.rule"))
            _set_line(line._element.spPr, ink, width_emu)
        elif kind == "band" and chrome_hex(ch, guide, "fill"):
            shape = _without_theme_style(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h)))
            shape.name = ch["chrome_id"]
            _solid_fill(shape._element.spPr, chrome_hex(ch, guide, "fill"))  # type: ignore[arg-type]
            _set_line(shape._element.spPr, None, None)
        elif kind == "logo":
            if ch.get("image_path"):
                picture_file(ch["image_path"], f"chrome {ch['chrome_id']} image_path")
                pic = slide.shapes.add_picture(ch["image_path"], Emu(x), Emu(y), Emu(w), Emu(h))
                pic.name = ch["chrome_id"]
            else:
                # no image bytes travel with a guide entry: a placeholder the agent can see
                # and replace, painted only in what the entry itself names
                fill_hex, line_hex = chrome_hex(ch, guide, "fill"), chrome_hex(ch, guide, "line")
                noted = notes is not None and any(n["chrome_id"] == ch["chrome_id"] for n in notes)  # once per deck, not per page
                if not fill_hex and not line_hex:
                    if notes is not None and not noted:
                        notes.append({"chrome_id": ch["chrome_id"], "kind": "logo", "drawn_as": None,
                                      "reason": "the entry carries no image_path, fill or line; add image_path, or a fill_hex/fill_role or line_hex/line_role (with line_width_pt) for a placeholder"})
                    continue
                shape = _without_theme_style(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h)))
                shape.name = ch["chrome_id"]
                if fill_hex:
                    _solid_fill(shape._element.spPr, fill_hex)
                else:
                    _insert_fill_first(shape._element.spPr, etree.Element(qn("a:noFill")))
                if line_hex and not ch.get("line_width_pt"):
                    guide.take(ch["chrome_id"], "logo placeholder outline 0.75pt (the entry records no line_width_pt)", look=True)
                _set_line(shape._element.spPr, line_hex, round(float(ch.get("line_width_pt") or 0.75) * EMU_PER_PT) if line_hex else None)
                if notes is not None and not noted:
                    notes.append({"chrome_id": ch["chrome_id"], "kind": "logo", "drawn_as": "rect placeholder", "fill_hex": fill_hex, "line_hex": line_hex,
                                  "reason": "the guide entry carries no image bytes; add image_path to the entry or place the picture by hand"})
        elif text is not None:
            rotation = float(ch.get("rotation_deg") or 0)
            box_w = max(w, 457200) if not rotation else w
            if not rotation and (ch.get("anchor") == "slide_right" or x + w > 0.8 * page_w):
                x = x + w - box_w  # keep the observed right edge: a wider box grows leftwards
            tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(box_w), Emu(max(h, 152400) if not rotation else h))
            tb.name = ch["chrome_id"]
            tf = tb.text_frame
            _body_props(tf._txBody, "t")
            tf._txBody.find(qn("a:bodyPr")).set("wrap", "none")  # the entry is one line of the reference: it never wraps
            p = tf.paragraphs[0]
            style = ch.get("text_style") or {}
            face = style.get("family") or guide.families.get("minor")
            if not face:
                raise SpecError(f"chrome {ch['chrome_id']}: its text_style names no family and the guide has no minor face; complete the entry's text_style")
            size_pt = float(style.get("size_pt") or guide.roles.get("type.caption", {}).get("size_pt") or 9)
            for key, value in (("family", f"{face} (the guide's minor face)"), ("size_pt", f"{size_pt}pt (type.caption's size, else 9pt)"), ("bold", "regular")):
                if style.get(key) is None:
                    guide.take(ch["chrome_id"], f"chrome text {key}: {value}; the entry's text_style records none", look=key == "size_pt" and not guide.roles.get("type.caption", {}).get("size_pt"))
            if text:
                run = p.add_run()
                # the space the reference sets between its words and the page number stays
                # ("Version 39", never "Version1")
                before = (ch.get("text") or "").split("<page number>")[0]
                run.text = text + (" " if "<page number>" in (ch.get("text") or "") and before[-1:].isspace() else "")
                run.font.name = face
                run.font.size = Pt(size_pt)
                run.font.bold = bool(style.get("bold"))
                if style.get("ink_hex"):
                    run.font.color.rgb = RGBColor.from_string(style["ink_hex"].upper())
            if kind == "page_number" or "<page number>" in (ch.get("text") or ""):
                fld = etree.SubElement(p._p, qn("a:fld"))
                fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF5B2F}")
                fld.set("type", "slidenum")
                rpr = etree.SubElement(fld, qn("a:rPr"))
                rpr.set("sz", str(int(size_pt * 100)))
                if style.get("ink_hex"):
                    sf = etree.SubElement(rpr, qn("a:solidFill"))
                    etree.SubElement(sf, qn("a:srgbClr")).set("val", style["ink_hex"].upper())
                etree.SubElement(rpr, qn("a:latin")).set("typeface", face)
                etree.SubElement(fld, qn("a:t")).text = "1"
            if not rotation and (ch.get("anchor") == "slide_right" or x + box_w > 0.8 * page_w):
                p.alignment = PP_ALIGN.RIGHT
        else:
            continue
        if ch.get("rotation_deg"):
            # box_in is the pre-rotation box: turned about its centre, the shape covers what the reference's does
            # python-pptx appends each new shape to the spTree: the last child is the one just drawn
            slide.shapes._spTree[-1].find(".//" + qn("a:xfrm")).set("rot", str(round(float(ch["rotation_deg"]) * 60000)))
        drawn.append(ch["chrome_id"])
        present.add(ch["chrome_id"])
    return drawn


def _picture_background(slide, path: str) -> None:
    """The slide's background as the reference's ground picture, stretched over the page
    (p:bg/p:bgPr/a:blipFill): the ground the reference paints under every page."""
    _, rid = slide.part.get_or_add_image_part(path)
    c_sld = slide._element.find(qn("p:cSld"))
    for old in c_sld.findall(qn("p:bg")):
        c_sld.remove(old)
    bg = etree.Element(qn("p:bg"))
    bg_pr = etree.SubElement(bg, qn("p:bgPr"))
    blip_fill = etree.SubElement(bg_pr, qn("a:blipFill"), rotWithShape="1")
    etree.SubElement(blip_fill, qn("a:blip")).set(qn("r:embed"), rid)
    etree.SubElement(etree.SubElement(blip_fill, qn("a:stretch")), qn("a:fillRect"))
    etree.SubElement(bg_pr, qn("a:effectLst"))
    c_sld.insert(0, bg)


def _send_to_back(slide, names: list[str]) -> None:
    """Drawn chrome goes behind the content and keeps its own order among itself (`names`
    is paint order): the spTree's first two children are its properties, so the first
    name lands at index 2, the next at 3."""
    tree = slide.shapes._spTree
    by_name: dict[str, etree._Element] = {}
    for el in list(tree):
        cnv = el.find(".//" + qn("p:cNvPr"))
        if cnv is not None and cnv.get("name") in names:
            by_name.setdefault(cnv.get("name"), el)
    for i, name in enumerate(n for n in names if n in by_name):
        tree.remove(by_name[name])
        tree.insert(2 + i, by_name[name])


def _origin_tag_xml() -> bytes:
    root = etree.Element(qn("p:tagLst"), nsmap={"p": P_NS})
    tag = etree.SubElement(root, qn("p:tag"))
    tag.set("name", ORIGIN_TAG_NAME)
    tag.set("val", ORIGIN_TAG_VALUE)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def write_origin_tag(slide) -> str:
    """One User Defined Tags part per generated slide, referenced from p:cSld/p:custDataLst
    (deckgen render/emit/chrome.py add_tag_part and attach_tag_reference; 2E-1 measured
    the part surviving a PowerPoint save with its p:tag intact). python-pptx writes the
    part's content type into [Content_Types].xml and the relationship into the slide's
    rels on save. p:custDataLst follows p:spTree in CT_CommonSlideData's sequence."""
    package = slide.part.package
    partname = package.next_partname("/ppt/tags/tag%d.xml")
    part = Part(partname, CT.PML_TAGS, package, _origin_tag_xml())
    rid = slide.part.relate_to(part, RT.TAGS)
    common = slide._element.find(qn("p:cSld"))
    if common is None:
        raise SpecError("the new slide has no p:cSld to carry its origin tag")
    custom = common.find(qn("p:custDataLst"))
    if custom is None:
        custom = etree.Element(qn("p:custDataLst"))
        common.find(qn("p:spTree")).addnext(custom)
    tags = etree.SubElement(custom, qn("p:tags"))
    tags.set(qn("r:id"), rid)
    return str(partname)


PLACEHOLDER_KIND = {"page_number": "SLIDE_NUMBER", "footer": "FOOTER", "date": "DATE"}


def _layers(ch: dict) -> list[str]:
    return sorted(set(ch.get("layers") or [ch.get("source_layer") or "slide"]))


# --- the reference firm's own chrome: keep, replace or drop ---------------------------------

#: A deck's `firm_chrome` choice per entry: "keep", "drop", {"text": "..."} or
#: {"image_path": "..."}.
FIRM_DECISIONS = ("keep", "drop", "replace")


def firm_chrome(guide_data: dict, mode: str, added: set[str] = frozenset()) -> list[dict]:
    """The chrome a generated page carries that can name the reference's firm: every entry
    chrome_on_page draws on the page that holds words (a label, or a footer the page
    actually draws, with two letters in a row; page numbers excluded) or a picture (kind
    logo). A deck for another firm, or for a made-up one, must not wear the reference's
    name: these are the entries the deck chooses for, by chrome_id, in `firm_chrome`. With
    no reference there is no reference firm: the chrome is the model's own."""
    import plan
    if not plan.has_reference(guide_data):
        return []
    page = (guide_data.get("reference") or {}).get("page") or {}
    pw, ph = page.get("width_in"), page.get("height_in")
    out = []
    for ch in guide_chrome(guide_data, mode) + [ch for cid, ch in minority_chrome(guide_data).items() if cid in added]:
        x, y, w, h = ch["box_in"]
        if pw is not None and ph is not None and not (x < pw and y < ph and x + w > 0 and y + h > 0):
            continue  # off the page: a note on the master's pasteboard, never seen
        layers = _layers(ch)
        if ch["kind"] == "footer" and mode != "draw" and "slide" not in layers:
            continue  # a layout's footer placeholder is drawn only when copied onto the slide
        words = re.search(r"[A-Za-z]{2}", chrome_text(ch) or "") is not None
        if ch["kind"] == "logo" or (ch["kind"] in ("label", "footer") and words):
            out.append(ch)
    return out


def firm_choices(spec: dict, candidates: list[dict]) -> dict[str, tuple[str, dict]]:
    """The deck's choices, proven at the boundary: each names a candidate once, with a
    decision in FIRM_DECISIONS; a replacement text for a text entry, an existing picture
    for a logo."""
    raw = spec.get("firm_chrome")
    if raw is None:
        return {}
    by_id = {ch["chrome_id"]: ch for ch in candidates}
    listing = ", ".join(f"{ch['chrome_id']} ({ch['kind']}{': ' + repr(chrome_text(ch)) if chrome_text(ch) else ''})" for ch in candidates) or "none"
    if not isinstance(raw, dict):
        raise SpecError(f"firm_chrome: an object {{chrome_id: \"keep\" | \"drop\" | {{\"text\": ...}} | {{\"image_path\": ...}}}}; the firm's chrome on this template: {listing}")
    out: dict[str, tuple[str, dict]] = {}
    for cid, choice in raw.items():
        ch = by_id.get(cid)
        if ch is None:
            raise SpecError(f"firm_chrome: {cid!r} is not chrome that names a firm on these pages; the entries are: {listing}")
        if choice in ("keep", "drop"):
            out[cid] = (choice, {})
        elif isinstance(choice, dict) and set(choice) == {"text"} and isinstance(choice["text"], str) and choice["text"].strip():
            if ch["kind"] == "logo":
                raise SpecError(f"firm_chrome {cid}: a logo is replaced by {{\"image_path\": ...}}, not text")
            out[cid] = ("replace", {"text": choice["text"].strip()})
        elif isinstance(choice, dict) and set(choice) == {"image_path"} and isinstance(choice["image_path"], str):
            if ch["kind"] != "logo":
                raise SpecError(f"firm_chrome {cid}: a {ch['kind']} is replaced by {{\"text\": ...}}, not a picture")
            ext = picture_file(choice["image_path"], f"firm_chrome {cid} image_path")
            out[cid] = ("replace", {"image_path": choice["image_path"], "format": ext})
        else:
            raise SpecError(f"firm_chrome {cid}: {choice!r} is not \"keep\", \"drop\", {{\"text\": \"...\"}} or {{\"image_path\": \"...\"}}")
    return out


def _norm_words(text: str | None) -> str:
    return " ".join((text or "").split())


def _template_shapes_of(prs, ch: dict) -> list[tuple]:
    """(container part, shape) on the template's masters and layouts that draw a chrome
    entry: same box within a twentieth of an inch, and the entry's words or its picture."""
    found = []
    tol = 0.05 * 914400
    want = [v * 914400 for v in ch["box_in"]]
    text = _norm_words(chrome_text(ch))
    containers = [m for m in prs.slide_masters] + [lay for m in prs.slide_masters for lay in m.slide_layouts]
    for container in containers:
        for shape in container.shapes:
            box = [shape.left, shape.top, shape.width, shape.height]
            if any(v is None for v in box) or any(abs(a - b) > tol for a, b in zip(box, want)):
                continue
            if ch["kind"] == "logo":
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    found.append((container, shape))
            elif shape.has_text_frame and text and _norm_words(shape.text_frame.text).startswith(text[:40]):
                found.append((container, shape))
    return found


def apply_firm_chrome(prs, choices: dict, candidates: list[dict]) -> list[dict]:
    """Every candidate with the decision taken for it. Inherited chrome (on a master or a
    layout) is changed in this deck's own copy of the template; slide-layer chrome is
    changed where draw_chrome paints it (firm_entries)."""
    report = []
    for ch in candidates:
        decision, detail = choices.get(ch["chrome_id"], ("keep", {}))
        entry = {"chrome_id": ch["chrome_id"], "kind": ch["kind"], "text": chrome_text(ch), "layers": _layers(ch),
                 "decision": decision if decision != "replace" else f"replace with {detail.get('text') or detail.get('image_path')!r}",
                 "chosen": ch["chrome_id"] in choices}
        if decision != "keep" and _layers(ch) != ["slide"]:
            shapes = _template_shapes_of(prs, ch)
            if not shapes:
                raise SpecError(f"firm_chrome {ch['chrome_id']}: no shape on the template's masters or layouts draws it at {ch['box_in']} in "
                                f"({entry['text'] or 'a picture'!r}); the template is not the reference's, so the entry cannot be changed there")
            for container, shape in shapes:
                if decision == "drop":
                    shape._element.getparent().remove(shape._element)
                elif "text" in detail:
                    paragraphs = shape.text_frame.paragraphs
                    runs = paragraphs[0].runs
                    if runs:
                        runs[0].text = detail["text"]
                        for r in runs[1:]:
                            r._r.getparent().remove(r._r)
                    else:
                        paragraphs[0].text = detail["text"]
                    for para in paragraphs[1:]:
                        para._p.getparent().remove(para._p)
                else:
                    blip = shape._element.find(".//" + qn("a:blip"))
                    image_part = container.part.related_part(blip.get(qn("r:embed")))
                    new = Path(detail["image_path"])
                    if detail["format"] != Path(str(image_part.partname)).suffix.lower().lstrip(".").replace("jpeg", "jpg"):
                        raise SpecError(f"firm_chrome {ch['chrome_id']}: the template's picture is {Path(str(image_part.partname)).suffix}; "
                                        f"{new} is {detail['format'].upper()}: give a picture of the same type")
                    image_part._blob = new.read_bytes()
            entry["changed_on"] = sorted({f"{type(c).__name__} {getattr(c, 'name', '')}".strip() for c, _ in shapes})
        report.append(entry)
    return report


def firm_entries(entries: list[dict], choices: dict) -> list[dict]:
    """Slide-layer chrome as the deck chose it: a dropped entry is not drawn, a replaced
    one draws the deck's text or picture."""
    out = []
    for ch in entries:
        decision, detail = choices.get(ch["chrome_id"], ("keep", {}))
        if decision == "drop":
            continue
        out.append({**ch, **detail} if decision == "replace" else ch)
    return out


def emit(spec: dict, template: str, guide_data: dict, out: str, chrome: str) -> dict:
    guide = Guide(guide_data)
    prs = Presentation(template)
    kept = len(prs.slides)  # a template made with --keep carries pages the new ones follow
    surface = guide.colors.get("color.surface", {}).get("hex")
    def furniture(page: dict | None) -> tuple[set[str], list[dict]]:
        """Under --chrome auto, the placeholders a page copies from its layout and the
        slide-layer entries it draws itself (furniture the reference pastes on every slide:
        a logo, a rule, a label; nothing inherits it from the layout, so it is drawn; an
        entry with a layout or master layer is inherited and never redrawn)."""
        wanted, own = set(), []
        if chrome == "auto":
            for ch in chrome_on_page(guide_data, chrome, page):
                layers = _layers(ch)
                if "slide" in layers and ch["kind"] in PLACEHOLDER_KIND:
                    wanted.add(PLACEHOLDER_KIND[ch["kind"]])
                if layers == ["slide"]:
                    own.append(ch)
        return wanted, firm_entries(own, choices)

    candidates = firm_chrome(guide_data, chrome, added_chrome(spec))
    choices = firm_choices(spec, candidates)
    wanted_chrome, slide_only = furniture(None)
    chrome_report: dict = {"mode": chrome, "placeholders_copied": sorted(wanted_chrome), "slide_layer_entries": [ch["chrome_id"] for ch in slide_only], "notes": [],
                           "firm": apply_firm_chrome(prs, choices, candidates)}
    report = {"out": out, "slides": [], "chrome": chrome_report}
    for si, sd in enumerate(spec["slides"], 1):
        guide.taken = []
        page_choices = page_chrome(guide_data, chrome, sd.get("chrome"), f"slide {si}")
        wanted_chrome, slide_only = furniture(page_choices)
        layout = _layout_named(prs, sd.get("layout"))
        slide = prs.slides.add_slide(layout)
        filled: set[int] = set()
        title = sd.get("title")
        if title:
            title_ph = next((ph for ph in slide.placeholders if "TITLE" in str(ph.placeholder_format.type)), None)
            if title_ph is not None and not title.get("box"):
                spec_t = guide.type_spec(title.get("role", "type.title"), where="title")
                title_ph.text_frame.text = ""
                p = title_ph.text_frame.paragraphs[0]
                run = p.add_run()
                run.text = apply_case(title["text"], spec_t.get("case") or "as_typed")
                ink = guide.hex(title.get("ink_role"), where="title") if title.get("ink_role") else (spec_t.get("ink") if HEX_RE.match(spec_t.get("ink") or "") else None)
                _style_run(run, spec_t, guide, emphasis=False, ink_hex=ink)
                if title.get("align"):
                    p.alignment = ALIGN.get(title["align"], PP_ALIGN.LEFT)
                filled.add(title_ph.placeholder_format.idx)
            else:
                box = tuple(int(v) for v in title["box"])
                write_text(slide, {"id": "title", "kind": "text", "paragraphs": [{"role": title.get("role", "type.title"), "runs": [{"text": title["text"]}], "lines": title.get("lines"), "fit": title.get("fit")}]}, guide, box)
        # a text element whose first paragraph is the title takes the layout's title
        # placeholder when the layout has one
        title_el = next((el for el in sd.get("elements", []) if el.get("kind") == "text" and el.get("paragraphs") and el["paragraphs"][0].get("role") == "type.title"), None)
        title_ph = None
        if title_el is not None and not title:
            title_ph = next((ph for ph in slide.placeholders if "TITLE" in str(ph.placeholder_format.type)), None)
            if title_ph is not None:
                filled.add(title_ph.placeholder_format.idx)
        for ph in list(slide.placeholders):
            if ph.placeholder_format.idx not in filled:
                ph._element.getparent().remove(ph._element)  # an unfilled prompt is not content
        charts: list[dict] = []
        for el in sd.get("elements", []):
            box = tuple(int(v) for v in el["box"])
            if any(v < 0 for v in box[2:]):
                raise SpecError(f"{el.get('id')}: negative extent")
            if el.get("foreign"):
                pass
            kind = el["kind"]
            if kind == "text":
                write_text(slide, el, guide, box, shape=title_ph if el is title_el else None)
                if el is title_el and title_ph is not None:
                    # the layout put the title placeholder first in the tree; it paints where
                    # the title stands in the page's order, over a panel filled before it
                    # (mckinsey p5: a white panel holding the title hid it; the contrast gate
                    # had read the title over the panel, as now drawn)
                    sp = title_ph._element
                    tree = sp.getparent()
                    tree.remove(sp)
                    ext = tree.find(qn("p:extLst"))
                    if ext is not None:
                        ext.addprevious(sp)
                    else:
                        tree.append(sp)
            elif kind == "rule":
                write_rule(slide, el, guide, box)
            elif kind == "rect":
                write_rect(slide, el, guide, box)
            elif kind == "image":
                write_image(slide, el, box, guide)
            elif kind == "table":
                write_table(slide, el, guide, box)
            elif kind == "chart":
                charts.append(write_chart(slide, el, guide, box))
            else:
                raise SpecError(f"{el.get('id')}: unknown element kind {kind!r}")
        if wanted_chrome:
            _copy_chrome_placeholders(slide, layout, wanted_chrome)
        drawn: list[str] = []
        page_w, page_h = int(prs.slide_width), int(prs.slide_height)
        if chrome == "draw":
            ground = (guide_data.get("ground") or {}).get("image_path")
            if ground and Path(ground).is_file():
                _picture_background(slide, ground)
            elif surface and surface.upper() != "FFFFFF":
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor.from_string(surface.upper())
            drawn = draw_chrome(slide, guide, guide_data, page_w, page_h, entries=firm_entries(chrome_on_page(guide_data, "draw", page_choices), choices),
                                notes=chrome_report["notes"])
        elif chrome == "auto" and slide_only:
            # a slide-layer page number, footer or date whose placeholder the layout carries
            # was copied above; the rest is drawn, each entry once, named by its chrome_id
            layout_kinds = {str(sp.placeholder_format.type).split(".")[-1].split(" ")[0] for sp in layout.placeholders}
            entries = [ch for ch in slide_only if PLACEHOLDER_KIND.get(ch["kind"]) not in layout_kinds]
            drawn = draw_chrome(slide, guide, guide_data, page_w, page_h, entries=entries, notes=chrome_report["notes"])
        if drawn:
            _send_to_back(slide, drawn)
        tag_part = write_origin_tag(slide)
        if sd.get("notes"):
            slide.notes_slide.notes_text_frame.text = sd["notes"]
        entry = {"slide": si, "layout": layout.name, "elements": len(sd.get("elements", [])), "chrome_drawn": drawn, "origin_tag": tag_part}
        if page_choices:
            entry["chrome_choices"] = page_choices
        if charts:
            entry["charts"] = charts
        if guide.taken:
            # every value written here that neither the spec nor the guide states
            entry["defaults_taken"] = {}
            for where, what in guide.taken:
                entry["defaults_taken"].setdefault(where, []).append(what)
        report["slides"].append(entry)
    prs.save(out)
    return report


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec")
    ap.add_argument("--template", required=True)
    ap.add_argument("--guide", required=True)
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--chrome", default="auto", choices=["auto", "none", "draw"], help="draw: paint the guide's chrome and surface on every slide (PDF references, blank templates)")
    args = ap.parse_args(argv)
    import cli
    cli.guard_args(args, "out")
    spec = json.loads(Path(args.spec).read_text())
    guide = json.loads(Path(args.guide).read_text())
    print(json.dumps(emit(spec, args.template, guide, args.out, args.chrome), indent=1))
    return 0


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import cli
    cli.run(main)
