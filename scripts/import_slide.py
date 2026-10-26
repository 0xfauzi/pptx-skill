#!/usr/bin/env python3
"""Import a reference slide into a page spec of roles, for the round trip.

    python3 scripts/import_slide.py ref.pptx --survey --guide guide.json           # which page to round-trip
    python3 scripts/import_slide.py ref.pptx --slide 3 --guide guide.json -o page.deck.json
    python3 scripts/build_deck.py page.deck.json --guide guide.json --template tpl.pptx -o rt.pptx --reference ref.pptx

The round trip is the one acceptance test the extractor cannot game: a slide the house
made, read into roles and weights (never its coordinates), solved with measured fit,
emitted on the reference's own master, must land near the house when the judge compares
it with the reference. It proves the extractor, the solver and the emitter agree. It is
also the seed for new pages: a spec imported from the exemplar of a composition family
starts inside the house's own bands and columns.

What survives the import: the title, every text shape's paragraphs with their bullet
levels, bold runs, fills and inverse ink, the row and column arrangement (shapes that
overlap vertically share a row, weights follow widths), pictures (a pptx picture with its
own bytes, a PDF picture cut from the page render) and a pptx's native tables (as table
specs, a column or row bold throughout as an emphasis group). What does not: charts (and
a PDF's tables) become a placeholder line to rebuild from data; turned text is left out;
merged cells are written unmerged. Each slide's `import_lost` lists these; `--survey`
lists them for every page, with whether each imported page fits, to choose the page.

A PDF page imports from ooxml.py's text blocks (`--slide` is the page number): the title
is the block at the deck's recurring title position, a block's fill is the filled
rectangle under its centre, text at a position where the guide mined chrome is left to
the chrome, and text turned a quarter is left out (a page spec sets text upright). Build
it on the blank template with --chrome draw.
"""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
from pathlib import Path

from PIL import Image

sys.dont_write_bytecode = True  # the skill directory is never written, not even a bytecode cache
sys.path.insert(0, str(Path(__file__).resolve().parent))

from extract_style import POSITION_DIAMETER_EMU  # noqa: E402  the grain chrome is mined at
from ooxml import Shape, SlidePrint, delta_e, fingerprint_deck, paragraphs_words  # noqa: E402

EMU_PER_INCH = 914400
CHROME_PH = {"sldNum", "ftr", "dt"}


def role_table(guide: dict) -> list[tuple[str, float, int, str | None]]:
    out = []
    for name, spec in guide["typography"]["roles"].items():
        if spec.get("size_pt"):
            out.append((name, float(spec["size_pt"]), int(spec.get("weight") or 400), spec.get("ink_role")))
    return out


PREFERENCE = ["type.body", "type.h2", "type.h1", "type.emphasis", "type.caption", "type.footnote", "type.source_line",
              "type.kicker", "type.table_cell", "type.table_header", "type.axis_label", "type.data_label", "type.title", "type.subtitle"]


def bind_role(size: float, bold: bool, is_title: bool, roles: list[tuple[str, float, int, str | None]]) -> str:
    if is_title and any(r[0] == "type.title" for r in roles):
        return "type.title"
    weight = 700 if bold else 400
    best = None
    for name, rsize, rweight, _ in roles:
        if name == "type.title":
            continue
        cost = abs(rsize - size) * 10 + (0 if rweight == weight else 3) + (PREFERENCE.index(name) if name in PREFERENCE else 20) * 0.01
        if best is None or cost < best[0]:
            best = (cost, name)
    return best[1] if best else "type.body"


def colour_roles(guide: dict) -> dict[str, str]:
    out = {}
    for name, spec in guide.get("color", {}).get("roles", {}).items():
        if spec.get("hex"):
            out[name] = spec["hex"].upper()
    for name, spec in guide.get("extra_roles", {}).items():
        if spec.get("hex"):
            out[name] = spec["hex"].upper()
    return out


def nearest_colour(hex_value: str, roles: dict[str, str], prefixes: tuple[str, ...] = ()) -> tuple[str | None, float]:
    """Nearest role by CIE76; `prefixes` keeps text ink away from surface roles of the same
    hex (white text on a blue box is ink_inverse, not surface)."""
    best = None
    for name, h in roles.items():
        if prefixes and not any(name.startswith(p) for p in prefixes):
            continue
        d = delta_e(hex_value.upper(), h)
        if best is None or d < best[1]:
            best = (name, d)
    return best if best else (None, 999.0)


def overlap_v(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> float:
    top = max(a[1], b[1])
    bottom = min(a[1] + a[3], b[1] + b[3])
    return max(0, bottom - top) / max(1, min(a[3], b[3]))


def rows_of(shapes: list[Shape]) -> list[list[Shape]]:
    """Shapes overlapping vertically by half their smaller height share a row."""
    rows: list[list[Shape]] = []
    for sh in sorted(shapes, key=lambda s: (s.box[1], s.box[0])):  # type: ignore[index]
        for row in rows:
            if any(overlap_v(sh.box, o.box) >= 0.5 for o in row):  # type: ignore[arg-type]
                row.append(sh)
                break
        else:
            rows.append([sh])
    for row in rows:
        row.sort(key=lambda s: s.box[0])  # type: ignore[index]
    return rows


def text_element(sh: Shape, ident: str, roles: list[tuple[str, float, int, str | None]], colours: dict[str, str], styling: dict, notes: list[str],
                 fill: str | None = None, extra_fills: tuple[str, ...] = ()) -> dict:
    fill = fill or sh.fill
    paragraphs = []
    para_style: dict[str, dict] = {}
    inverse = None
    for i, para in enumerate(sh.paragraphs or []):
        runs: list[dict] = []
        for r in para["runs"]:
            if r.get("br") or not r["text"].strip():
                if runs:  # a break between two words is a space between them, never nothing
                    runs[-1]["text"] += " "
                continue
            runs.append(dict(r, text=r["text"].replace("\n", " ")))
        if not runs:
            continue
        lead = runs[0]
        all_bold = all(r["bold"] for r in runs)
        role = bind_role(lead["size_pt"], all_bold, sh.ph_type == "title", roles)
        text = paragraphs_words([para])  # explicit breaks and runs of spaces collapse
        p = {"role": role, "text": text}
        if para["bullet"].startswith("char") or para["bullet"] == "auto":
            p["bullet_level"] = int(para["level"])
        paragraphs.append(p)
        if not all_bold and any(r["bold"] for r in runs):
            para_style[str(len(paragraphs) - 1)] = {"runs": [{"text": r["text"], **({"emphasis": True} if r["bold"] else {})} for r in runs]}
        if para.get("align") in ("ctr", "r"):
            para_style.setdefault(str(len(paragraphs) - 1), {})["align"] = {"ctr": "center", "r": "right"}[para["align"]]
        if lead.get("color"):
            name, d = nearest_colour(lead["color"], colours, ("color.ink", "color.accent", "color.rule"))
            # the ink is written whenever it differs from the bound role's own ink: a role
            # bound by size can carry another ink (a white box header) than the run had
            own = next((r[3] for r in roles if r[0] == role), None) or "color.ink"
            if name and d < 12 and name != own:
                para_style.setdefault(str(len(paragraphs) - 1), {})["ink_role"] = name
                if name == "color.ink_inverse":
                    inverse = name
    el = {"type": "element", "id": ident, "kind": "text", "paragraphs": paragraphs}
    st: dict = {}
    if sh.text_anchor in ("ctr", "b"):
        st["anchor"] = "m" if sh.text_anchor == "ctr" else "b"
    if fill:
        # a fill may take any role a fill can name: the surfaces, accents, rule and the extra roles
        name, d = nearest_colour(fill, colours, ("color.surface", "color.accent", "color.rule", *extra_fills))
        if name and d < 12:
            st["fill_role"] = name
            st["pad_role"] = "space.xs"
        else:
            notes.append(f"{ident}: fill {fill} has no colour role within CIE76 12 (nearest {name} at {d:.1f}); left unfilled")
    if para_style:
        st["paragraphs"] = para_style
    if st:
        styling[ident] = st
    return el


CHROME_TEXT_KINDS = {"label", "page_number", "footer", "date"}


def chrome_boxes(guide: dict) -> list[tuple[int, int, int, int]]:
    """The guide's text chrome positions in EMU. A PDF page carries its footer and page
    number as page text, so the slide's own copy is recognised by where the guide found
    that chrome (a band or a rule is never a text block's position)."""
    out = []
    for c in guide.get("chrome") or []:
        box = c.get("box_in")
        if c.get("kind") in CHROME_TEXT_KINDS and isinstance(box, list) and len(box) == 4:
            out.append(tuple(int(v * EMU_PER_INCH) for v in box))  # type: ignore[arg-type]
    return out  # type: ignore[return-value]


def is_guide_chrome(box: tuple[int, int, int, int], chrome: list[tuple[int, int, int, int]]) -> bool:
    return any(abs(box[0] - cx) <= POSITION_DIAMETER_EMU and abs(box[1] - cy) <= POSITION_DIAMETER_EMU for cx, cy, _, _ in chrome)


def fill_under(sh: Shape, s: SlidePrint) -> str | None:
    """A PDF paints a text box's fill as its own rectangle under the text: the fill of the
    smallest filled rectangle holding the block's centre stands in for the shape fill a
    pptx declares. The page's own ground is no fill."""
    x, y, w, h = sh.box  # type: ignore[misc]
    cx, cy = x + w // 2, y + h // 2  # the centre: a line can overhang its box by a glyph (bain p45 tab)
    best: tuple[int, str] | None = None
    for o in s.shapes:
        if o.kind != "shape" or not o.fill or not o.box:
            continue
        ox, oy, ow, oh = o.box
        if ox <= cx <= ox + ow and oy <= cy <= oy + oh and (best is None or ow * oh < best[0]):
            best = (ow * oh, o.fill)
    return best[1] if best and best[1] != s.background else None


PICTURE_DPI = 150  # the page render a PDF picture is cut from; the build scales it into the solved box


def pdf_picture(path: str, index: int, box: tuple[int, int, int, int], out_dir: Path, ident: str) -> str | None:
    """A PDF picture's pixels: the page rendered once by pdftoppm and the picture's box cut
    from it, so the round trip draws the reference's own image, not an empty frame."""
    page_png = out_dir / f"{Path(path).stem}-page{index + 1}.png"
    if not page_png.exists():
        subprocess.run(["pdftoppm", "-png", "-r", str(PICTURE_DPI), "-f", str(index + 1), "-l", str(index + 1), "-singlefile", path,
                        str(page_png.with_suffix(""))], check=True, timeout=300)
    scale = PICTURE_DPI / EMU_PER_INCH
    with Image.open(page_png) as page:
        left, top = max(0, round(box[0] * scale)), max(0, round(box[1] * scale))
        right, bottom = min(page.width, round((box[0] + box[2]) * scale)), min(page.height, round((box[1] + box[3]) * scale))
        if right - left < 2 or bottom - top < 2:
            return None
        target = out_dir / f"{Path(path).stem}-page{index + 1}-{ident}.png"
        page.crop((left, top, right, bottom)).save(target, format="PNG")
    return str(target)


def _pptx_shape(path: str, index: int, ref: str):
    """The python-pptx shape a fingerprint ref names (`pic[id=3]`, `graphicFrame[id=5]`)
    on slide `index`, searched through groups; None when the ref names no id."""
    import re

    from pptx import Presentation
    m = re.search(r"\[id=(\d+)\]", ref or "")
    if not m:
        return None
    want = int(m.group(1))

    def walk(shapes):
        for sh in shapes:
            if sh.shape_id == want:
                return sh
            if sh.shape_type is not None and getattr(sh, "shapes", None) is not None:
                found = walk(sh.shapes)
                if found is not None:
                    return found
        return None

    return walk(Presentation(path).slides[index].shapes)


def pptx_picture(path: str, index: int, ref: str, out_dir: Path, ident: str) -> str | None:
    """The picture's own bytes, written beside the imported spec: the house pastes its
    exhibits as pictures, and a round trip without them draws empty frames."""
    sh = _pptx_shape(path, index, ref)
    if sh is None or not hasattr(sh, "image"):
        return None
    out = out_dir / f"rt-s{index + 1}-{ident}.{sh.image.ext}"
    out.write_bytes(sh.image.blob)
    return str(out)


def pptx_table(path: str, index: int, ref: str, ident: str, notes: list[str], lost: list[str]) -> dict | None:
    """A native table as a table spec: the first row its header, every other row a body
    row, the cell text as written; a column or row whose every body cell is bold is an
    emphasis group that keeps its fill. Fills, rules and merged cells are the guide's
    table look, not the reference's: the notes say so."""
    sh = _pptx_shape(path, index, ref)
    if sh is None or not getattr(sh, "has_table", False):
        return None
    grid = [[cell for cell in row.cells] for row in sh.table.rows]
    if len(grid) < 2:
        return None
    text = [[" ".join(c.text_frame.text.split()) for c in row] for row in grid]
    # a cell's weight as written on its runs: True, False, or None for an empty cell, which
    # says nothing about its column's or row's weight
    bold = [[None if not c.text_frame.text.strip() else all(bool(r.font.bold) for p in c.text_frame.paragraphs for r in p.runs if r.text.strip())
             for c in row] for row in grid]
    columns = [{"key": f"c{i}", "label": t or " ", "value_kind": "text"} for i, t in enumerate(text[0])]
    body, body_bold = text[1:], bold[1:]
    def all_bold(cells: list) -> bool:
        written = [c for c in cells if c is not None]
        return bool(written) and all(written)

    every = all_bold([c for row in body_bold for c in row])  # a table bold throughout marks nothing
    emphasis = [] if every else [{"columns": [f"c{i}"], "fill_role": "keep"} for i in range(len(columns)) if all_bold([row[i] for row in body_bold])]
    emphasis += [] if every else [{"rows": [r], "fill_role": "keep"} for r, row in enumerate(body_bold) if all_bold(row)]
    merged = sum(1 for row in grid for c in row if c.is_merge_origin or c.is_spanned)
    notes.append(f"{ident}: native table {len(grid)} x {len(columns)} imported as a table spec (header row 1); its fills and rules are the "
                 f"guide's table look, not this table's own" + (f"; {merged} merged cells written unmerged" if merged else ""))
    if merged:
        lost.append(f"table {ident}'s {merged} merged cells (written unmerged)")
    el = {"type": "element", "id": ident, "kind": "table", "table": {"columns": columns, "rows": [[c or " " for c in row] for row in body]}}
    if emphasis:
        el["emphasis"] = emphasis
    return el


def import_slide(path: str, index: int, guide: dict, image_dir: Path | None = None) -> dict:
    slides = fingerprint_deck(path)  # every page: a PDF title is found from the position titles recur at
    if not 0 <= index < len(slides):
        raise SystemExit(f"{path}: slide {index + 1} does not exist ({len(slides)} slides)")
    s = slides[index]
    roles = role_table(guide)
    colours = colour_roles(guide)
    extra = tuple(name for name, spec in (guide.get("extra_roles") or {}).items() if spec.get("hex"))
    notes: list[str] = []
    lost: list[str] = []  # what the build cannot draw from this page, in words: the round trip's own loss list
    styling: dict = {}
    images: dict = {}
    chrome = chrome_boxes(guide) if s.source == "pdf" else []
    content = []
    for sh in s.shapes:
        if not sh.box or sh.ph_type in CHROME_PH or sh.kind in ("line",) or not (sh.text_chars or sh.kind in ("picture", "table", "chart")):
            continue
        if chrome and is_guide_chrome(sh.box, chrome):
            notes.append(f"{sh.ref}: {sh.text!r} sits at a chrome position of the guide; the build draws chrome, so it is not imported")
            continue
        if sh.rotation:
            notes.append(f"{sh.ref}: text turned {sh.rotation // 60000} degrees ({sh.text!r}); a page spec sets text upright, so it is not imported")
            lost.append(f"turned text {sh.text!r}")
            continue
        content.append(sh)
    title = next((sh for sh in content if sh.ph_type == "title"), None)
    body = [sh for sh in content if sh is not title]
    children: list[dict] = []
    if title is not None:
        children.append(text_element(title, "title", roles, colours, styling, notes, fill_under(title, s) if s.source == "pdf" else None, extra))
    n = 0
    for row in rows_of(body):
        elems = []
        for sh in row:
            n += 1
            ident = f"e{n}"
            if sh.kind == "picture":
                elems.append({"type": "element", "id": ident, "kind": "image", "size": {"width_emu": sh.box[2], "height_emu": sh.box[3]}})  # type: ignore[index]
                if image_dir is None:
                    crop = None
                elif s.source == "pdf":
                    crop = pdf_picture(path, index, sh.box, image_dir, ident)  # type: ignore[arg-type]
                else:
                    crop = pptx_picture(path, index, sh.ref, image_dir, ident)
                if crop:
                    images[ident] = {"path": crop}
                    notes.append(f"{ident}: picture {sh.box[2] / EMU_PER_INCH:.2f} x {sh.box[3] / EMU_PER_INCH:.2f} in, "  # type: ignore[index]
                                 f"{'cut from the page render' if s.source == 'pdf' else 'its own bytes'}: {crop}")
                else:
                    notes.append(f"{ident}: picture {sh.box[2] / EMU_PER_INCH:.2f} x {sh.box[3] / EMU_PER_INCH:.2f} in, bytes not imported")  # type: ignore[index]
                    lost.append(f"picture {ident}'s bytes (an empty frame)")
            elif sh.kind == "table" and s.source != "pdf" and (table := pptx_table(path, index, sh.ref, ident, notes, lost)) is not None:
                elems.append(table)
            elif sh.kind in ("table", "chart"):
                elems.append({"type": "element", "id": ident, "kind": "text", "paragraphs": [{"role": "type.body", "text": f"[{sh.kind} to rebuild from data]"}]})
                notes.append(f"{ident}: {sh.kind} left as a placeholder; rebuild it from the dataset")
                lost.append(f"{sh.kind} {ident} (a placeholder line; rebuild it from data)")
            else:
                elems.append(text_element(sh, ident, roles, colours, styling, notes, fill_under(sh, s) if s.source == "pdf" else None, extra))
        if len(elems) == 1:
            children.append(elems[0])
        else:
            total = sum(sh.box[2] for sh in row)  # type: ignore[index]
            for el, sh in zip(elems, row):
                el["weight"] = round(sh.box[2] / total, 4)  # type: ignore[index]
            # a stretched row is flexible, so the body's rows share the content height by
            # weight instead of packing at the top (the reference's label boxes are taller
            # than their text and sit at one pitch down the page)
            children.append({"type": "region", "id": f"row{len(children)}", "layout": {"kind": "flex", "direction": "row", "gap_role": "space.gutter",
                                                                                          "pad_role": None, "align": "stretch", "justify": "start"},
                             "children": elems})
    # the house fills its page: the body rows share the content height (weights), as the
    # reference's rows do (keybank slide 3: seven rows at one pitch down to the footer rule)
    title_el = children[0] if children and children[0].get("id") == "title" else None
    body_children = children[1:] if title_el else children
    for row in body_children:
        if row.get("type") == "region" or row.get("kind") == "text":
            row["weight"] = 1.0
    body = {"type": "region", "id": "body", "weight": 1.0, "layout": {"kind": "flex", "direction": "column", "gap_role": "space.sm",
                                                                     "pad_role": None, "align": "stretch", "justify": "start"},
            "children": body_children}
    root = {"type": "region", "id": "root", "layout": {"kind": "flex", "direction": "column", "gap_role": "space.sm", "pad_role": "space.margin",
                                                        "align": "stretch", "justify": "start"},
            "children": ([title_el] if title_el else []) + [body]}
    slide = {"id": f"import-{index + 1}", "layout": s.layout_name or "Title and Content",
             "notes": f"round trip of {Path(path).name} slide {index + 1}: roles and weights only, no coordinates",
             "spec": {"root": root, "splittable_groups": []}, "styling": styling, "import_notes": notes,
             "import_lost": lost + ["each element's exact position and size (rebuilt from roles and weights)",
                                    "each table's own fills, rules and column widths (the guide's table look; widths measured)"
                                    if any('"kind": "table"' in json.dumps(c) for c in body_children) else None]}
    unknown = json.dumps(root, ensure_ascii=False).count("�")
    if unknown:
        # ooxml reads a glyph the PDF maps to no character as U+FFFD: the words are not known
        slide["import_lost"].insert(0, f"{unknown} character(s) the PDF maps to no character (written as �): retype them from the render")
    slide["import_lost"] = [x for x in slide["import_lost"] if x]
    if images:
        slide["images"] = images
    return {"slides": [slide]}


def survey(path: str, guide: dict, chrome: str = "auto") -> list[dict]:
    """Every page of the reference as the round trip would rebuild it: what it holds, what
    the build cannot draw from it (import_lost), and whether the imported page solves at
    full scale. The page to round-trip is a content page the build can draw: one with a
    title, text and an exhibit, losing least, that fits."""
    import tempfile

    import component_api
    import layout as L
    from build_deck import solve_slide
    lg = L.load_guide(guide)
    ruler, index = L.default_ruler(), L.default_index()
    api = component_api.make_api(guide, lg, ruler, index)
    content = set((((guide.get("variety") or {}).get("content_pages")) or {}).get("pages") or [])
    triage = {p["slide"]: p.get("triage") for p in (guide.get("variety") or {}).get("pages") or []}
    out = []
    with tempfile.TemporaryDirectory() as tmp:
        for i in range(len(fingerprint_deck(path))):
            sl = import_slide(path, i, guide, Path(tmp))["slides"][0]
            kinds: dict[str, int] = {}
            for node in json.dumps(sl["spec"]).split('"kind": "')[1:]:
                k = node.split('"', 1)[0]
                if k in ("text", "table", "image", "chart", "rule"):
                    kinds[k] = kinds.get(k, 0) + 1
            try:
                _, _, outcome = solve_slide(json.loads(json.dumps(sl)), guide, chrome, lg, ruler, index, api, L.Fraction(1))
                fits = outcome.status
            except (SystemExit, L.LayoutContractError) as error:
                fits = f"refused: {error}"[:160]
            has_title = '"id": "title"' in json.dumps(sl["spec"])
            exhibit = kinds.get("table", 0) + kinds.get("image", 0)
            lost = [x for x in sl["import_lost"] if not x.startswith(("each element", "each table"))]
            page = i + 1
            out.append({"slide": page, "layout": sl["layout"], "content": (page in content) if content else triage.get(page) == "content",
                        "title": has_title, "text": kinds.get("text", 0) - int(has_title), "exhibits": exhibit, "fits_full_scale": fits,
                        "lost": lost})
    for p in out:
        p["candidate"] = bool(p["content"] and p["title"] and p["text"] > 0 and p["exhibits"] > 0 and not p["lost"] and p["fits_full_scale"] == "fit")
    return out


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("reference")
    ap.add_argument("--slide", type=int, help="1-based slide number")
    ap.add_argument("--survey", action="store_true", help="list every page: what it holds, what the round trip would lose, whether it fits")
    ap.add_argument("--guide", required=True, help="the completed guide (roles and colours to bind to)")
    ap.add_argument("--chrome", default="auto", choices=["auto", "none", "draw"], help="with --survey: as the build will run")
    ap.add_argument("-o", "--out")
    args = ap.parse_args(argv)
    import cli
    cli.guard_args(args, "out")
    guide = json.loads(cli.in_path(args.guide, "--guide").read_text())
    cli.in_path(args.reference, "reference")
    if args.survey:
        pages = survey(args.reference, guide, args.chrome)
        layouts: dict[str, int] = {}
        for p in pages:
            if p["content"]:
                layouts[p["layout"]] = layouts.get(p["layout"], 0) + 1
        common = max(layouts, key=layouts.get) if layouts else None
        for p in pages:
            mark = "CANDIDATE" if p["candidate"] and p["layout"] == common else ("candidate" if p["candidate"] else "")
            print(f"slide {p['slide']:>3} {mark:<9} {p['layout']!s:<24} content={p['content']!s:<5} title={p['title']!s:<5} text={p['text']} "
                  f"exhibits={p['exhibits']} fits={p['fits_full_scale']}" + (f"  loses: {'; '.join(p['lost'])}" if p["lost"] else ""))
        best = [p["slide"] for p in pages if p["candidate"] and p["layout"] == common] or [p["slide"] for p in pages if p["candidate"]]
        print(f"round-trip candidates (content, title, text and an exhibit, nothing lost, fits; on the house's most used layout {common!r} first): "
              f"{best or 'none: pick the content page that loses least and name every loss in the report'}")
        if args.out:
            Path(args.out).write_text(json.dumps(pages, indent=1))
        return 0
    if args.slide is None or not args.out:
        ap.error("--slide and -o are required (or --survey)")
    Path(args.out).resolve().parent.mkdir(parents=True, exist_ok=True)  # a new work directory, as build_deck makes for -o
    deck = import_slide(args.reference, args.slide - 1, guide, Path(args.out).resolve().parent)
    Path(args.out).write_text(json.dumps(deck, indent=1))
    sl = deck["slides"][0]
    print(json.dumps({"out": args.out, "layout": sl["layout"], "elements": sum(1 for _ in json.dumps(sl["spec"]).split('"kind":')) - 1,
                      "styled": sorted(sl["styling"]), "lost": sl["import_lost"], "notes": sl["import_notes"]}, indent=1))
    return 0


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import cli
    cli.run(main)
