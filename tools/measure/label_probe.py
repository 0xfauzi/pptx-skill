#!/usr/bin/env python3
"""Measure where LibreOffice sets chart data labels, and where their ink sits in the word box,
per face and size.

    python3 -I tools/measure/label_probe.py OUT_DIR [--faces Arial,Calibri,Georgia,Inter] [--sizes 7,9,12]

One page per face and size, written through emit_pptx.write_chart (the build's emitter), four
charts a page:
  1. a line chart with no markers, 12 points, every label `above` its point;
  2. the same with every label `below`;
  3. a horizontal clustered bar chart, 4 bars, labels `inside_base`;
  4. the same with labels `inside_end`.
Each label alternates a figure with no descender ("12.5") and a word with one ("Aug"), so the
ink band is read for both. The deck is rendered with scripts/render.py; pdftotext -bbox
gives every word's box, which is matched to the point or bar it labels (chart_geometry
predicts both). Reported per face and size, in EMU and as a share of the size:

- above: the gap from the bottom of the word box to its point; below: from the point to
  the top of the word box (build_deck.LINE_LABEL_GAP_EMU);
- box height over the size (the 1.21 the build uses);
- inside_base: the gap from the bar's base to the word box's left; inside_end: from the
  word box's right to the bar's end (build_deck.BAR_LABEL_INSET_EMU);
- ink: the first and last rows holding ink inside each word box, rasterised at 600 dpi, as
  shares of the box height from its top and its bottom, split by descender
  (build_deck.INK_TOP_INSET, INK_BOTTOM_INSET).

The faces are named as a deck names them; fontconfig maps each to the face the renderer
draws (Liberation Sans, Carlito, Gelasio), and the report prints that mapping. Nothing
here is a unit test: it is the measurement those constants are quoted from.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

SKILL = Path(__file__).resolve().parents[2] / "plugins" / "designing-presentations" / "skills" / "designing-presentations"
sys.path.insert(0, str(SKILL / "scripts"))
import emit_pptx as E
from chart_geometry import ChartGeometry

PAGE_W, PAGE_H = 12_192_000, 6_858_000
EMU_PER_PT = 12_700
COLOURS = {"color.ink": "000000", "color.line": "C8C8C8", "color.bar": "F2D27A", "color.axis": "808080", "color.surface": "FFFFFF"}
VALUES = [20, 70, 30, 85, 25, 60, 15, 90, 35, 75, 40, 65]
BARS = [92, 74, 55, 38]
DESCENDERS = set("gjpqy,;()[]{}|/")


def guide(face: str, size: float) -> dict:
    return {"color": {"roles": {k: {"hex": v} for k, v in COLOURS.items()}},
            "typography": {"families": {"major": face, "minor": face},
                           "roles": {"type.data_label": {"size_pt": size, "family": "minor"}, "type.axis_label": {"size_pt": size, "family": "minor"}}},
            "spacing": {"steps": {}}, "tables": {}}


def text_of(i: int, v: float) -> str:
    return "Aug" if i % 2 else f"{v:.1f}"


def charts(face: str, size: float) -> list[dict]:
    hidden = {"visible": False}
    cat = {"visible": False, "line_role": "none", "tick_marks": "none", "reverse": False, "gridlines": "none", "labels": {"show": False}}
    line_axes = {"category": cat, "value": {**hidden, "min": 0, "max": 100, "cross_between": "between"}}
    bar_axes = {"category": {**cat, "reverse": True}, "value": {**hidden, "min": 0, "max": 100, "cross_between": "between"}}
    lab = {"show": True, "role": "type.data_label", "ink_role": "color.ink"}
    out = []
    w, h = (PAGE_W - 3 * 457_200) // 2, (PAGE_H - 3 * 457_200) // 2
    boxes = [(457_200, 457_200), (457_200 * 2 + w, 457_200), (457_200, 457_200 * 2 + h), (457_200 * 2 + w, 457_200 * 2 + h)]
    for k, pos in enumerate(("above", "below")):
        out.append({"id": f"line-{pos}", "kind": "chart", "type": "LINE", "categories": [f"c{i}" for i in range(len(VALUES))], "number_format": "0.0",
                    "series": [{"name": "s", "values": VALUES, "line_role": "color.line", "line_width_pt": 0.75, "smooth": False, "marker": {"style": "none"},
                                "labels": {**lab, "position": pos}, "points": [{"index": i, "label": {"text": text_of(i, v)}} for i, v in enumerate(VALUES)]}],
                    "axes": line_axes, "legend": "none", "title": "none", "plot_fill_role": "none", "text": {"role": "type.axis_label", "ink_role": "color.ink"},
                    "plot_area": {"x": 0.05, "y": 0.15, "w": 0.9, "h": 0.7}, "box": [*boxes[k], w, h]})
    for k, pos in enumerate(("inside_base", "inside_end")):
        out.append({"id": f"bar-{pos}", "kind": "chart", "type": "BAR_CLUSTERED", "categories": [f"b{i}" for i in range(len(BARS))], "number_format": "0.0",
                    "series": [{"name": "s", "values": BARS, "fill_role": "color.bar", "line_role": "none",
                                "labels": {**lab, "position": pos}, "points": [{"index": i, "label": {"text": text_of(i, v)}} for i, v in enumerate(BARS)]}],
                    "axes": bar_axes, "gap_width": 60, "overlap": 0, "legend": "none", "title": "none", "plot_fill_role": "none",
                    "text": {"role": "type.axis_label", "ink_role": "color.ink"}, "plot_area": {"x": 0.05, "y": 0.05, "w": 0.9, "h": 0.9},
                    "box": [*boxes[2 + k], w, h]})
    return out


def words(pdf: Path, page: int) -> list[tuple[str, float, float, float, float]]:
    html = subprocess.run(["pdftotext", "-f", str(page), "-l", str(page), "-bbox", str(pdf), "-"], capture_output=True, text=True, check=True).stdout
    return [(m.group(5), float(m.group(1)), float(m.group(2)), float(m.group(3)), float(m.group(4)))
            for m in re.finditer(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">([^<]*)</word>', html)]


def ink_rows(img: np.ndarray, box_pt: tuple, dpi: int) -> tuple[float, float] | None:
    """The top and bottom of the ink inside a word box, in points: the first and last pixel
    rows whose darkest pixel is under half grey."""
    s = dpi / 72
    x0, y0, x1, y1 = (round(box_pt[0] * s), int(box_pt[1] * s), round(box_pt[2] * s), int(np.ceil(box_pt[3] * s)))
    rows = np.where(img[y0:y1, x0:x1].min(axis=1) < 128)[0]
    if not len(rows):
        return None
    return (y0 + rows[0]) / s, (y0 + rows[-1] + 1) / s


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("out")
    ap.add_argument("--faces", default="Arial,Calibri,Georgia,Inter")
    ap.add_argument("--sizes", default="7,9,12")
    ap.add_argument("--dpi", type=int, default=600)
    args = ap.parse_args(argv)
    out = Path(args.out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    faces = args.faces.split(",")
    sizes = [float(s) for s in args.sizes.split(",")]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(PAGE_W), Emu(PAGE_H)
    pages = []
    for face in faces:
        for size in sizes:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            g = E.Guide(guide(face, size))
            els = charts(face, size)
            for el in els:
                E.write_chart(slide, el, g, tuple(el["box"]))
            pages.append((face, size, els))
    deck = out / "labels.pptx"
    prs.save(deck)
    pdf = out / "labels.pdf"
    subprocess.run([sys.executable, "-I", str(SKILL / "scripts" / "render.py"), str(deck), "--out", str(out / "render"), "--keep-pdf", str(pdf)],
                   capture_output=True, text=True, timeout=1800, check=True)
    mapped = {f: subprocess.run(["fc-match", "-f", "%{family[0]}", f], capture_output=True, text=True, check=True).stdout for f in faces}
    report: dict = {"renderer_faces": mapped, "pages": []}
    for n, (face, size, els) in enumerate(pages, 1):
        subprocess.run(["pdftoppm", "-f", str(n), "-l", str(n), "-r", str(args.dpi), "-gray", "-png", str(pdf), str(out / f"p{n}")], check=True)
        img = np.asarray(Image.open(next(out.glob(f"p{n}-*.png"))).convert("L"), dtype=float)
        ws = words(pdf, n)
        row: dict = {"face": face, "renderer_face": mapped[face], "size_pt": size, "above": [], "below": [], "box_over_size": [],
                     "inside_base": [], "inside_end": [], "ink_top": {"descender": [], "none": []}, "ink_bottom": {"descender": [], "none": []}}
        for el in els:
            geo = ChartGeometry(el, tuple(el["box"]))
            pos = el["series"][0]["labels"]["position"]
            for i, v in enumerate(el["series"][0]["values"]):
                text = text_of(i, v)
                if el["type"] == "LINE":
                    px_, py_ = (c / EMU_PER_PT for c in geo.point(0, i))
                    near = [w for w in ws if w[0] == text and abs((w[1] + w[3]) / 2 - px_) < 6]
                    if not near:
                        continue
                    w = min(near, key=lambda w: abs((w[2] + w[4]) / 2 - py_))
                    row[pos].append(round(((py_ - w[4]) if pos == "above" else (w[2] - py_)) * EMU_PER_PT))
                else:
                    x, y, bw, bh = (c / EMU_PER_PT for c in geo.bar_rect(0, i))
                    near = [w for w in ws if w[0] == text and y <= (w[2] + w[4]) / 2 <= y + bh and x - 1 <= (w[1] + w[3]) / 2 <= x + bw + 1]
                    if not near:
                        continue
                    w = near[0]
                    row[pos].append(round(((w[1] - x) if pos == "inside_base" else (x + bw - w[3])) * EMU_PER_PT))
                row["box_over_size"].append(round((w[4] - w[2]) / size, 3))
                ink = ink_rows(img, w[1:], args.dpi)
                if ink is not None:
                    kind = "descender" if set(text) & DESCENDERS else "none"
                    hgt = w[4] - w[2]
                    row["ink_top"][kind].append(round((ink[0] - w[2]) / hgt, 3))
                    row["ink_bottom"][kind].append(round((w[4] - ink[1]) / hgt, 3))
        report["pages"].append(row)
        print(f"{face:8s} ({mapped[face]}) {size:g}pt: above {min(row['above'], default=None)}..{max(row['above'], default=None)} "
              f"below {min(row['below'], default=None)}..{max(row['below'], default=None)} EMU; box/size {min(row['box_over_size'], default=None)}.."
              f"{max(row['box_over_size'], default=None)}; inside_base {row['inside_base']} inside_end {row['inside_end']} EMU; "
              f"ink top min {min(row['ink_top']['none'] + row['ink_top']['descender'], default=None)}, bottom min "
              f"none {min(row['ink_bottom']['none'], default=None)} descender {min(row['ink_bottom']['descender'], default=None)}")
    (out / "labels.json").write_text(json.dumps(report, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
